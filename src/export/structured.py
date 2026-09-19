"""结构化流式 DOCX 导出。

普通正文按逻辑块写入 Word 段落；表格写成 ``w:tbl``，公式优先写成
OMML，图片写成内嵌图片。该路径依赖 Document IR 的阅读顺序，避免把每个
字符或每个文本片段都变成绝对定位对象。
"""

from __future__ import annotations

from html.parser import HTMLParser
from io import BytesIO
from dataclasses import dataclass
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..fonts.embedding import FontPlan, attach_embedded_fonts
from ..ir.model import IRBlock, IRDocument, IRPage, IRTextLine
from ..ooxml_positioning import (
    add_absolute_picture,
    add_absolute_shape,
    add_absolute_text_box,
    detach_header_footer,
    new_canvas_paragraph,
    position_table,
)
from ..page_render import (
    _add_bookmark,
    _add_formula_omml_element,
    _add_toc_field,
    _bookmark_name,
)
from .document_setup import (
    StageCallback,
    _notify_stage,
    _remove_initial_empty_paragraph,
    _set_document_styles,
    _set_section_page,
    render_page_image_png,
    render_pdf_region,
)
from .table_render import _add_pdf_table


STRUCTURED_MODES = frozenset({"structured", "flow"})
FLOW_MARGIN_POINTS = 46.8
DENSE_VECTOR_LIMIT = 64
FLOW_TEXT_KINDS = frozenset(
    {
        "title",
        "heading1",
        "heading2",
        "heading3",
        "paragraph",
        "body",
        "line",
        "ordered",
        "bullet",
        "plugin",
        "caption",
        "quote",
        "code",
    }
)
_HEADING_LEVELS = {"title": 1, "heading1": 1, "heading2": 2, "heading3": 3}
_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


@dataclass(frozen=True)
class _HtmlTableCell:
    text: str
    row_span: int = 1
    column_span: int = 1
    header: bool = False


class _HtmlTableParser(HTMLParser):
    """提取 OCR 表格 HTML 中的行和单元格文本。"""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[_HtmlTableCell]] = []
        self._row: list[_HtmlTableCell] | None = None
        self._cell: list[str] | None = None
        self._cell_row_span = 1
        self._cell_column_span = 1
        self._cell_header = False

    def handle_starttag(
        self,
        tag: str,
        attrs: list[tuple[str, str | None]],
    ) -> None:
        if tag.lower() == "tr":
            self._row = []
        elif tag.lower() in {"td", "th"} and self._row is not None:
            self._cell = []
            attributes = {
                str(name).lower(): str(value or "")
                for name, value in attrs
            }
            self._cell_row_span = self._positive_span(attributes.get("rowspan"))
            self._cell_column_span = self._positive_span(
                attributes.get("colspan")
            )
            self._cell_header = tag.lower() == "th"

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"td", "th"}:
            if self._row is not None and self._cell is not None:
                value = " ".join("".join(self._cell).split())
                self._row.append(
                    _HtmlTableCell(
                        text=value,
                        row_span=self._cell_row_span,
                        column_span=self._cell_column_span,
                        header=self._cell_header,
                    )
                )
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self._row:
                self.rows.append(self._row)
            self._row = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    @staticmethod
    def _positive_span(value: str | None) -> int:
        try:
            return max(int(value or "1"), 1)
        except ValueError:
            return 1


def _html_table_rows(value: str) -> list[list[_HtmlTableCell]]:
    parser = _HtmlTableParser()
    parser.feed(value)
    return parser.rows


def _set_run_style(
    run: Any,
    *,
    font_name: str = "",
    font_size: float = 0.0,
    bold: bool = False,
    italic: bool = False,
    color: tuple[int, int, int] | None = None,
) -> None:
    run.font.name = font_name or "Microsoft YaHei"
    if font_size > 0:
        run.font.size = Pt(font_size)
    run.bold = bool(bold)
    run.italic = bool(italic)
    if color is not None and len(color) >= 3:
        run.font.color.rgb = RGBColor(
            int(color[0]) & 0xFF,
            int(color[1]) & 0xFF,
            int(color[2]) & 0xFF,
        )
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is not None:
        fonts.set(qn("w:eastAsia"), font_name or "Microsoft YaHei")


def _planned_font(
    font_plan: FontPlan | None,
    *,
    raw_name: str,
    fallback_name: str,
    bold: bool,
    italic: bool,
) -> tuple[str, bool, bool]:
    if font_plan is None:
        return fallback_name or "Microsoft YaHei", bool(bold), bool(italic)
    name = font_plan.word_name_for(raw_name, fallback_name)
    planned_bold, planned_italic = font_plan.style_for(
        raw_name,
        bool(bold),
        bool(italic),
    )
    return name or fallback_name or "Microsoft YaHei", planned_bold, planned_italic


def _configure_paragraph(paragraph: Any, block: IRBlock | None = None) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if block is None:
        return
    alignment = str(block.alignment or "").lower()
    if alignment in _ALIGNMENTS:
        paragraph.alignment = _ALIGNMENTS[alignment]
    # The layout annotator stores line_spacing as a multiple of the source
    # font size (for example 1.10), not as an absolute point value.  Passing
    # that multiple directly to Pt() produces a roughly 1-point exact line
    # height and makes every wrapped line overlap.
    if block.line_spacing > 0 and len(block.lines) > 1:
        font_size = float(block.font_size or 0.0)
        if font_size > 0:
            paragraph.paragraph_format.line_spacing = Pt(
                max(font_size * float(block.line_spacing), font_size)
            )
    if block.first_line_indent > 0:
        paragraph.paragraph_format.first_line_indent = Pt(block.first_line_indent)


def _line_spans_match(line: IRTextLine, block: IRBlock) -> bool:
    spans = list(line.spans or ())
    return bool(
        spans
        and len(block.lines) == 1
        and block.text == line.text
        and "".join(span.text for span in spans) == line.text
    )


def _add_text_runs(
    paragraph: Any,
    block: IRBlock,
    *,
    font_plan: FontPlan | None = None,
) -> None:
    if len(block.lines) == 1 and _line_spans_match(block.lines[0], block):
        for span in block.lines[0].spans:
            font_name, bold, italic = _planned_font(
                font_plan,
                raw_name=str(getattr(span, "pdf_font_name", "") or ""),
                fallback_name=str(getattr(span, "font_name", "") or ""),
                bold=bool(getattr(span, "bold", False)),
                italic=bool(getattr(span, "italic", False)),
            )
            run = paragraph.add_run(span.text)
            _set_run_style(
                run,
                font_name=font_name,
                font_size=float(getattr(span, "font_size", 0.0) or 0.0),
                bold=bold,
                italic=italic,
                color=getattr(span, "color", None),
            )
        return
    text = block.text or "\n".join(line.text for line in block.lines)
    font_name, bold, italic = _planned_font(
        font_plan,
        raw_name=str(getattr(block, "pdf_font_name", "") or ""),
        fallback_name=str(getattr(block, "font_name", "") or ""),
        bold=bool(getattr(block, "bold", False)),
        italic=bool(getattr(block, "italic", False)),
    )
    run = paragraph.add_run(text)
    _set_run_style(
        run,
        font_name=font_name,
        font_size=float(block.font_size or 0.0),
        bold=bold,
        italic=italic,
        color=block.color,
    )


def _add_text_block(
    document: Any,
    block: IRBlock,
    *,
    font_plan: FontPlan | None = None,
) -> Any:
    level = _HEADING_LEVELS.get(block.kind)
    if level is not None:
        paragraph = document.add_heading(level=level)
        _configure_paragraph(paragraph, block)
        _add_text_runs(paragraph, block, font_plan=font_plan)
        return paragraph
    if block.kind == "ordered":
        paragraph = document.add_paragraph(style="List Number")
    elif block.kind in {"bullet", "plugin"}:
        paragraph = document.add_paragraph(style="List Bullet")
    elif block.kind == "code":
        paragraph = document.add_paragraph(style="Code Block")
    else:
        paragraph = document.add_paragraph()
    _configure_paragraph(paragraph, block)
    _add_text_runs(paragraph, block, font_plan=font_plan)
    return paragraph


def _image_size_points(
    block: IRBlock,
    page: IRPage,
    *,
    margin_points: float,
) -> tuple[float, float]:
    if block.bbox is not None:
        width = max(float(block.bbox[2]) - float(block.bbox[0]), 1.0)
        height = max(float(block.bbox[3]) - float(block.bbox[1]), 1.0)
    else:
        width = max(float(block.image_width or 0.0), 1.0)
        height = max(float(block.image_height or 0.0), 1.0)
    max_width = max(page.width - margin_points * 2.0, 24.0)
    max_height = max(page.height - margin_points * 2.0, 24.0)
    scale = min(1.0, max_width / width, max_height / height)
    return width * scale, height * scale


def _image_alignment(block: IRBlock, page: IRPage) -> int:
    if block.bbox is None:
        return WD_ALIGN_PARAGRAPH.CENTER
    center = (float(block.bbox[0]) + float(block.bbox[2])) / 2.0
    if abs(center - page.width / 2.0) <= max(page.width * 0.12, 24.0):
        return WD_ALIGN_PARAGRAPH.CENTER
    if center > page.width * 0.58:
        return WD_ALIGN_PARAGRAPH.RIGHT
    return WD_ALIGN_PARAGRAPH.LEFT


def _add_inline_image(
    container: Any,
    block: IRBlock,
    page: IRPage,
    *,
    margin_points: float,
    paragraph: Any | None = None,
) -> Any | None:
    if not block.image_bytes:
        return None
    created_paragraph = paragraph is None
    target = paragraph or container.add_paragraph()
    target.alignment = _image_alignment(block, page)
    _configure_paragraph(target)
    width, height = _image_size_points(
        block,
        page,
        margin_points=margin_points,
    )
    try:
        inline = target.add_run().add_picture(
            BytesIO(block.image_bytes),
            width=Inches(width / 72.0),
            height=Inches(height / 72.0),
        )
    except Exception:
        if created_paragraph:
            element = target._element
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)
        return None
    description = str(block.image_alt or block.source or "PDF 图片")
    inline._inline.docPr.set("descr", description)
    return target


def _render_block_region(
    source_pdf: Path | None,
    page: IRPage,
    block: IRBlock,
    *,
    max_pixels: int,
) -> bytes | None:
    if block.fallback_image:
        return block.fallback_image
    if source_pdf is None or block.bbox is None or not source_pdf.is_file():
        return None
    try:
        return render_pdf_region(
            source_pdf,
            page.page_number - 1,
            block.bbox,
            max_pixels=max_pixels,
            jpeg_quality=88,
            scale=3.0,
        )
    except Exception:
        return None


def _table_cell_line_count(value: Any) -> int:
    return sum(bool(line.strip()) for line in str(value or "").splitlines())


def _table_has_overfull_cells(table: Any) -> bool:
    """识别表格解析把多条物理行压进同一个单元格的情况。"""
    cells = list(getattr(table, "cells", ()) or ())
    if not cells or int(getattr(table, "column_count", 0) or 0) < 2:
        return False
    for cell in cells:
        line_count = _table_cell_line_count(getattr(cell, "text", ""))
        if line_count < 2:
            continue
        bbox = getattr(cell, "bbox", None)
        if not bbox or len(bbox) < 4:
            continue
        cell_height = max(float(bbox[3]) - float(bbox[1]), 0.0)
        font_size = max(float(getattr(cell, "font_size", 0.0) or 0.0), 6.0)
        required_height = line_count * font_size * 1.05
        if required_height > cell_height + 4.0:
            return True
    return False


def _horizontal_overlap_ratio(
    left: tuple[float, float, float, float],
    right: tuple[float, float, float, float],
) -> float:
    overlap = max(min(left[2], right[2]) - max(left[0], right[0]), 0.0)
    denominator = max(min(left[2] - left[0], right[2] - right[0]), 1.0)
    return overlap / denominator


def _complex_table_regions(page: IRPage) -> tuple[tuple[float, float, float, float], ...]:
    """合并同一复杂表格被边框检测拆开的相邻区域。"""
    table_blocks = [
        block
        for block in page.body_blocks
        if block.kind == "table"
        and block.bbox is not None
        and _table_has_overfull_cells(block.table)
    ]
    if not table_blocks:
        return ()

    all_table_blocks = [
        block
        for block in page.body_blocks
        if block.kind == "table" and block.bbox is not None
    ]
    regions: list[tuple[float, float, float, float]] = []
    used: set[int] = set()
    for seed in table_blocks:
        seed_index = all_table_blocks.index(seed)
        if seed_index in used:
            continue
        selected = [seed]
        used.add(seed_index)
        changed = True
        while changed:
            changed = False
            current_region = (
                min(float(block.bbox[0]) for block in selected if block.bbox),
                min(float(block.bbox[1]) for block in selected if block.bbox),
                max(float(block.bbox[2]) for block in selected if block.bbox),
                max(float(block.bbox[3]) for block in selected if block.bbox),
            )
            for index, candidate in enumerate(all_table_blocks):
                if index in used or candidate.bbox is None:
                    continue
                candidate_bbox = tuple(float(value) for value in candidate.bbox)
                vertical_gap = max(
                    candidate_bbox[1] - current_region[3],
                    current_region[1] - candidate_bbox[3],
                    0.0,
                )
                if (
                    _horizontal_overlap_ratio(current_region, candidate_bbox)
                    >= 0.8
                    and vertical_gap <= 120.0
                ):
                    selected.append(candidate)
                    used.add(index)
                    changed = True
        x0 = max(min(float(block.bbox[0]) for block in selected if block.bbox) - 2.0, 0.0)
        top = max(min(float(block.bbox[1]) for block in selected if block.bbox) - 2.0, 0.0)
        x1 = min(max(float(block.bbox[2]) for block in selected if block.bbox) + 2.0, page.width)
        bottom = min(max(float(block.bbox[3]) for block in selected if block.bbox) + 2.0, page.height)
        regions.append((x0, top, x1, bottom))
    return tuple(sorted(regions, key=lambda bbox: (bbox[1], bbox[0])))


def _block_overlaps_region(
    block: IRBlock,
    region: tuple[float, float, float, float],
) -> bool:
    if block.bbox is None:
        return False
    x0, top, x1, bottom = (float(value) for value in block.bbox)
    rx0, rtop, rx1, rbottom = region
    overlap_x = min(x1, rx1) - max(x0, rx0)
    center_y = (top + bottom) / 2.0
    return overlap_x > 0.0 and rtop <= center_y <= rbottom


def _is_formula_block_candidate(block: IRBlock) -> bool:
    if block.bbox is None:
        return False
    text = str(block.text or "").strip()
    if not text or len(text) > 90:
        return False
    marks = len(re.findall(r"[=≤≥≈+\-*/^_()]", text))
    return block.kind == "formula" and marks >= 2


def _is_compact_formula_fragment(block: IRBlock) -> bool:
    if block.bbox is None:
        return False
    text = str(block.text or "").strip()
    if not text or text.endswith((".", "。", ":", "：")):
        return False
    if block.kind == "formula":
        return _is_formula_block_candidate(block) or len(text) <= 24
    if len(text) > 24:
        return False
    if block.kind not in FLOW_TEXT_KINDS:
        return False
    if _is_formula_block_candidate(block):
        return True
    if len(text) > 10 or re.search(r"\b(?:the|and|of|with|from|where)\b", text, re.I):
        return False
    return bool(
        re.fullmatch(r"[A-Za-z0-9α-ωΑ-Ω₀-₉⁰-⁹\s()\[\]{}_^+\-*/=]+", text)
    )


def _formula_regions(page: IRPage) -> tuple[tuple[float, float, float, float], ...]:
    """合并公式及其分数线/上下标碎片，按源 PDF 区域转图。"""
    blocks = [block for block in page.body_blocks if block.bbox is not None]
    seeds = [block for block in blocks if _is_formula_block_candidate(block)]
    regions: list[tuple[float, float, float, float]] = []
    used: set[int] = set()
    for seed in seeds:
        seed_index = blocks.index(seed)
        if seed_index in used:
            continue
        selected = [seed]
        selected_indices = {seed_index}
        changed = True
        while changed:
            changed = False
            region = (
                min(float(block.bbox[0]) for block in selected if block.bbox),
                min(float(block.bbox[1]) for block in selected if block.bbox),
                max(float(block.bbox[2]) for block in selected if block.bbox),
                max(float(block.bbox[3]) for block in selected if block.bbox),
            )
            for index, candidate in enumerate(blocks):
                if index in selected_indices or not _is_compact_formula_fragment(candidate):
                    continue
                candidate_bbox = tuple(float(value) for value in candidate.bbox)
                vertical_gap = max(
                    candidate_bbox[1] - region[3],
                    region[1] - candidate_bbox[3],
                    0.0,
                )
                horizontal_gap = max(
                    candidate_bbox[0] - region[2],
                    region[0] - candidate_bbox[2],
                    0.0,
                )
                vertical_overlap = min(candidate_bbox[3], region[3]) - max(
                    candidate_bbox[1], region[1]
                )
                if (
                    vertical_gap <= 14.0
                    and (vertical_overlap > 0.0 or horizontal_gap <= 12.0)
                ):
                    selected.append(candidate)
                    selected_indices.add(index)
                    changed = True
        if len(selected) < 2:
            continue
        used.update(selected_indices)
        regions.append(
            (
                max(min(float(block.bbox[0]) for block in selected if block.bbox) - 4.0, 0.0),
                max(min(float(block.bbox[1]) for block in selected if block.bbox) - 4.0, 0.0),
                min(max(float(block.bbox[2]) for block in selected if block.bbox) + 4.0, page.width),
                min(max(float(block.bbox[3]) for block in selected if block.bbox) + 4.0, page.height),
            )
        )
    return tuple(sorted(regions, key=lambda bbox: (bbox[1], bbox[0])))


def _should_use_page_image_fallback(
    page: IRPage,
    *,
    source_pdf: Path | None,
) -> bool:
    """为旋转图表页启用页级保真图，避免逐词拼装破坏图形。"""
    if source_pdf is None or not source_pdf.is_file():
        return False
    rotated_count = sum(
        abs(float(block.rotation or 0.0)) >= 1.0
        for block in page.body_blocks
        if block.kind in FLOW_TEXT_KINDS
    )
    vector_count = sum(block.kind == "vector" for block in page.body_blocks)
    return rotated_count >= 8 and vector_count >= DENSE_VECTOR_LIMIT


def _add_page_image_fallback(
    document: Any,
    page: IRPage,
    *,
    source_pdf: Path | None,
    max_fallback_pixels: int,
    report: dict[str, Any],
) -> Any | None:
    """把无法可靠流式重建的复杂页面作为单个页面锚定图像。"""
    if source_pdf is None:
        return None
    try:
        image_bytes = render_page_image_png(
            source_pdf,
            page.page_number - 1,
            dpi=144.0,
            max_pixels=max_fallback_pixels,
        )
    except Exception:
        return None
    paragraph = new_canvas_paragraph(document)
    add_absolute_picture(
        document,
        image_bytes,
        x_points=0.0,
        y_points=0.0,
        width_points=max(page.width, 0.5),
        height_points=max(page.height, 0.5),
        z_order=0,
        object_id=2000 + len(report["placements"]),
        name=f"page_image_{page.page_number}",
        behind_text=True,
        paragraph=paragraph,
    )
    report["page_image_fallback_count"] += 1
    report["visual_only_page_count"] += 1
    return paragraph


def _vector_union_bbox(page: IRPage) -> tuple[float, float, float, float] | None:
    boxes = [
        block.bbox
        for block in page.body_blocks
        if block.kind == "vector" and block.bbox is not None
    ]
    if not boxes:
        return None
    return (
        max(min(float(box[0]) for box in boxes), 0.0),
        max(min(float(box[1]) for box in boxes), 0.0),
        min(max(float(box[2]) for box in boxes), page.width),
        min(max(float(box[3]) for box in boxes), page.height),
    )


def _dense_vector_is_form_background(
    page: IRPage,
    bbox: tuple[float, float, float, float],
) -> bool:
    """识别覆盖整页的表单矢量层，避免与可编辑正文重复栅格化。"""
    page_area = max(float(page.width) * float(page.height), 1.0)
    bbox_area = max(float(bbox[2] - bbox[0]), 0.0) * max(
        float(bbox[3] - bbox[1]),
        0.0,
    )
    text_count = sum(
        block.kind in FLOW_TEXT_KINDS and bool(str(block.text or "").strip())
        for block in page.body_blocks
    )
    return text_count >= 20 and bbox_area / page_area >= 0.70


def _dense_form_background_blocks(page: IRPage) -> tuple[IRBlock, ...]:
    """筛出密集表单中可独立重建的彩色大面积填充块。

    这类块只承担底色。筛选条件限制为无描边、非白色且横向覆盖页面大部分
    的矢量，避免把复选框和灰色字段块误当成页面背景。
    """
    selected: list[IRBlock] = []
    minimum_width = max(float(page.width) * 0.75, 1.0)
    for block in page.body_blocks:
        if block.kind != "vector" or block.bbox is None:
            continue
        vector = block.vector
        if vector is None:
            continue
        fill = getattr(vector, "fill_color", None)
        if fill is None or getattr(vector, "stroke_color", None) is not None:
            continue
        if all(int(channel) >= 245 for channel in fill[:3]):
            continue
        x0, top, x1, bottom = (float(value) for value in block.bbox)
        if x1 - x0 < minimum_width or bottom <= top:
            continue
        selected.append(block)
    return tuple(
        sorted(
            selected,
            key=lambda block: (
                float(block.bbox[1]) if block.bbox is not None else 0.0,
                float(block.bbox[0]) if block.bbox is not None else 0.0,
            ),
        )
    )


def _add_formula_block(
    document: Any,
    page: IRPage,
    block: IRBlock,
    *,
    source_pdf: Path | None,
    margin_points: float,
    max_fallback_pixels: int,
    report: dict[str, Any],
) -> Any:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _configure_paragraph(paragraph)
    formula_text = str(block.text or block.latex or "").strip()
    if formula_text and _add_formula_omml_element(paragraph, formula_text):
        report["formula_omml_count"] += 1
        return paragraph

    image_bytes = _render_block_region(
        source_pdf,
        page,
        block,
        max_pixels=max_fallback_pixels,
    )
    if image_bytes:
        image_block = IRBlock(
            kind="image",
            page=block.page,
            bbox=block.bbox,
            image_bytes=image_bytes,
            image_width=block.width,
            image_height=block.height,
            image_alt="公式图片兜底",
            source="formula_fallback",
        )
        image_paragraph = _add_inline_image(
            document,
            image_block,
            page,
            margin_points=margin_points,
            paragraph=paragraph,
        )
        if image_paragraph is not None:
            report["formula_image_fallback_count"] += 1
            return paragraph

    run = paragraph.add_run(formula_text)
    _set_run_style(
        run,
        font_name="Cambria Math",
        font_size=float(block.font_size or 10.0),
        italic=True,
    )
    report["formula_text_fallback_count"] += 1
    return paragraph


def _add_html_table(document: Any, value: str) -> Any | None:
    rows = _html_table_rows(value)
    if not rows:
        return None

    placements: list[tuple[int, int, _HtmlTableCell]] = []
    occupied: set[tuple[int, int]] = set()
    column_count = 0
    for row_index, values in enumerate(rows):
        column_index = 0
        for cell in values:
            while any(
                (row_index + row_offset, column_index + column_offset)
                in occupied
                for row_offset in range(cell.row_span)
                for column_offset in range(cell.column_span)
            ):
                column_index += 1
            placements.append((row_index, column_index, cell))
            for row_offset in range(cell.row_span):
                for column_offset in range(cell.column_span):
                    occupied.add(
                        (row_index + row_offset, column_index + column_offset)
                    )
            column_index += cell.column_span
            column_count = max(column_count, column_index)
    if column_count < 1:
        return None

    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, column_index, source_cell in placements:
        target_cell = table.cell(row_index, column_index)
        if source_cell.row_span > 1 or source_cell.column_span > 1:
            target_cell = target_cell.merge(
                table.cell(
                    row_index + source_cell.row_span - 1,
                    column_index + source_cell.column_span - 1,
                )
            )
        target_cell.text = source_cell.text
        for paragraph in target_cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            if source_cell.header:
                for run in paragraph.runs:
                    run.bold = True
    return table


def _remove_tables_added_after(document: Any, table_count: int) -> None:
    for table in list(document.tables)[table_count:]:
        element = table._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)


def _add_table_block(
    document: Any,
    page: IRPage,
    block: IRBlock,
    *,
    source_pdf: Path | None,
    margin_points: float,
    max_fallback_pixels: int,
    report: dict[str, Any],
    positioned: bool = False,
) -> Any | None:
    table = None
    existing_table_count = len(document.tables)
    try:
        if block.kind == "table" and block.table is not None:
            table = _add_pdf_table(
                document,
                block.table,
                bordered=bool(getattr(block.table, "has_borders", True)),
                exact_widths=positioned,
                cell_vertical_alignment="top",
                tight_cell_margins=positioned,
                cell_padding_points=(0.0, 0.0) if positioned else None,
                merge_cells=True,
                empty_text=positioned,
            )
        elif block.kind == "html_table":
            table = _add_html_table(document, block.text)
    except Exception as error:
        _remove_tables_added_after(document, existing_table_count)
        report["table_native_error_count"] += 1
        report["table_native_error_types"].append(type(error).__name__)
    if table is not None:
        report["table_count"] += 1
        if block.kind == "html_table":
            report["html_table_count"] += 1
        return table

    image_bytes = _render_block_region(
        source_pdf,
        page,
        block,
        max_pixels=max_fallback_pixels,
    )
    if image_bytes:
        image_block = IRBlock(
            kind="image",
            page=block.page,
            bbox=block.bbox,
            image_bytes=image_bytes,
            image_width=block.width,
            image_height=block.height,
            image_alt="表格图片兜底",
            source="table_fallback",
        )
        image_paragraph = _add_inline_image(
            document,
            image_block,
            page,
            margin_points=margin_points,
        )
        if image_paragraph is not None:
            report["table_image_fallback_count"] += 1
            return None
    report["skipped_block_count"] += 1
    return None


def _add_flow_header_footer(
    section: Any,
    page: IRPage,
    *,
    font_plan: FontPlan | None,
    report: dict[str, Any],
) -> None:
    for layer, blocks in (
        ("header", page.header_blocks),
        ("footer", page.footer_blocks),
    ):
        if not blocks:
            continue
        container = section.header if layer == "header" else section.footer
        container.is_linked_to_previous = False
        paragraphs = list(container.paragraphs)
        first = paragraphs[0] if paragraphs else container.add_paragraph()
        first.text = ""
        for index, block in enumerate(sorted(blocks, key=lambda item: item.z_order)):
            paragraph = first if index == 0 else container.add_paragraph()
            if block.kind in FLOW_TEXT_KINDS:
                _configure_paragraph(paragraph, block)
                _add_text_runs(paragraph, block, font_plan=font_plan)
                report["header_footer_count"] += 1
            elif block.kind == "image":
                if _add_inline_image(
                    container,
                    block,
                    page,
                    margin_points=FLOW_MARGIN_POINTS,
                    paragraph=paragraph,
                ):
                    report["header_footer_count"] += 1


def _record_page_report(
    report: dict[str, Any],
    page: IRPage,
    *,
    start_index: int,
) -> None:
    placements = report["placements"][start_index:]
    report["pages"].append(
        {
            "page": page.page_number,
            "route": page.route,
            "text_paragraph_count": sum(
                item["kind"] in FLOW_TEXT_KINDS and item["status"] == "native"
                for item in placements
            ),
            "table_count": sum(
                item["kind"] in {"table", "html_table"}
                and item["status"] == "native"
                for item in placements
            ),
            "image_count": sum(
                item["kind"] in {"image", "page_image"}
                and item["status"] == "native"
                for item in placements
            ),
            "placements": placements,
        }
    )


def _record_placement(
    report: dict[str, Any],
    block: IRBlock,
    *,
    status: str,
    reason: str = "",
) -> None:
    report["placements"].append(
        {
            "kind": block.kind,
            "page": block.page,
            "block_id": block.block_id,
            "status": status,
            "reason": reason,
            "source_char_ids": list(block.meta.get("source_char_ids", ())),
            "source_text": block.text,
        }
    )


def _add_page_block(
    document: Any,
    page: IRPage,
    block: IRBlock,
    *,
    source_pdf: Path | None,
    font_plan: FontPlan | None,
    margin_points: float,
    max_fallback_pixels: int,
    report: dict[str, Any],
) -> Any | None:
    if block.kind in FLOW_TEXT_KINDS:
        paragraph = _add_text_block(document, block, font_plan=font_plan)
        report["text_paragraph_count"] += 1
        _record_placement(report, block, status="native", reason="flow_paragraph")
        return paragraph
    if block.kind in {"table", "html_table"}:
        fallback_count = report["table_image_fallback_count"]
        table = _add_table_block(
            document,
            page,
            block,
            source_pdf=source_pdf,
            margin_points=margin_points,
            max_fallback_pixels=max_fallback_pixels,
            report=report,
        )
        _record_placement(
            report,
            block,
            status=(
                "native"
                if table is not None
                else (
                    "image_fallback"
                    if report["table_image_fallback_count"] > fallback_count
                    else "skipped"
                )
            ),
            reason=(
                "flow_table"
                if table is not None
                else (
                    "table_image_fallback"
                    if report["table_image_fallback_count"] > fallback_count
                    else "table_unavailable"
                )
            ),
        )
        return None
    if block.kind == "formula":
        omml_count = report["formula_omml_count"]
        image_fallback_count = report["formula_image_fallback_count"]
        text_fallback_count = report["formula_text_fallback_count"]
        paragraph = _add_formula_block(
            document,
            page,
            block,
            source_pdf=source_pdf,
            margin_points=margin_points,
            max_fallback_pixels=max_fallback_pixels,
            report=report,
        )
        report["formula_count"] += 1
        _record_placement(
            report,
            block,
            status=(
                "native"
                if report["formula_omml_count"] > omml_count
                else (
                    "image_fallback"
                    if report["formula_image_fallback_count"] > image_fallback_count
                    else "text_fallback"
                )
            ),
            reason=(
                "flow_formula_omml"
                if report["formula_omml_count"] > omml_count
                else (
                    "formula_image_fallback"
                    if report["formula_image_fallback_count"] > image_fallback_count
                    else "formula_text_fallback"
                )
            ),
        )
        return paragraph
    if block.kind in {"image", "page_image"}:
        paragraph = _add_inline_image(
            document,
            block,
            page,
            margin_points=margin_points,
        )
        if paragraph is None:
            report["skipped_block_count"] += 1
            _record_placement(report, block, status="skipped", reason="image_data_missing")
            return None
        report["image_count"] += 1
        if block.kind == "page_image":
            report["page_image_count"] += 1
        _record_placement(report, block, status="native", reason="inline_picture")
        return paragraph
    if block.kind == "vector":
        image_bytes = _render_block_region(
            source_pdf,
            page,
            block,
            max_pixels=max_fallback_pixels,
        )
        if image_bytes:
            image_block = IRBlock(
                kind="image",
                page=block.page,
                bbox=block.bbox,
                image_bytes=image_bytes,
                image_width=block.width,
                image_height=block.height,
                image_alt="PDF 矢量区域",
                source="vector_fallback",
            )
            image_paragraph = _add_inline_image(
                document,
                image_block,
                page,
                margin_points=margin_points,
            )
            if image_paragraph is not None:
                report["vector_image_fallback_count"] += 1
                _record_placement(
                    report,
                    block,
                    status="image_fallback",
                    reason="vector_rasterized",
                )
                return None
        report["skipped_block_count"] += 1
        _record_placement(report, block, status="skipped", reason="vector_unavailable")
        return None
    report["skipped_block_count"] += 1
    _record_placement(report, block, status="skipped", reason="unsupported_kind")
    return None


def _add_positioned_text_block(
    document: Any,
    page: IRPage,
    block: IRBlock,
    *,
    font_plan: FontPlan | None,
    canvas_paragraph: Any,
    report: dict[str, Any],
) -> Any | None:
    """把表单页的小文本块按源坐标写成可编辑的浮动段落。"""
    if block.bbox is None:
        return _add_page_block(
            document,
            page,
            block,
            source_pdf=None,
            font_plan=font_plan,
            margin_points=0.0,
            max_fallback_pixels=1,
            report=report,
        )
    lines = tuple(block.lines or ())
    if not lines:
        lines = (
            IRTextLine(
                text=str(block.text or ""),
                bbox=tuple(float(value) for value in block.bbox),
                font_name=str(getattr(block, "font_name", "") or ""),
                font_size=float(getattr(block, "font_size", 0.0) or 8.0),
                color=block.color,
                bold=bool(getattr(block, "bold", False)),
                italic=bool(getattr(block, "italic", False)),
                alignment=str(getattr(block, "alignment", "") or "left"),
                pdf_font_name=str(getattr(block, "pdf_font_name", "") or ""),
            ),
        )
    first_paragraph = None
    for line_index, line in enumerate(lines):
        x0, top, x1, bottom = (float(value) for value in line.bbox)
        width = max(x1 - x0, 1.0)
        line_font_size = float(line.font_size or block.font_size or 8.0)
        height = max(bottom - top, line_font_size, 1.0)
        font_name, bold, italic = _planned_font(
            font_plan,
            raw_name=str(getattr(line, "pdf_font_name", "") or ""),
            fallback_name=str(getattr(line, "font_name", "") or ""),
            bold=bool(getattr(line, "bold", False)),
            italic=bool(getattr(line, "italic", False)),
        )
        paragraph = add_absolute_text_box(
            document,
            x_points=x0,
            y_points=top,
            width_points=width,
            height_points=height,
            text=str(line.text or ""),
            font_name=font_name,
            font_size=line_font_size,
            bold=bold,
            italic=italic,
            color=getattr(line, "color", None) or block.color,
            alignment=str(getattr(line, "alignment", "") or block.alignment or "left"),
            z_order=int(getattr(line, "z_order", 0) or 0),
            object_id=5000 + len(report["placements"]) * 1000 + line_index,
            name=f"positioned_text_{block.block_id or 'block'}",
            paragraph=canvas_paragraph,
        )
        if first_paragraph is None:
            first_paragraph = paragraph
    report["text_paragraph_count"] += len(lines)
    _record_placement(report, block, status="native", reason="positioned_text")
    return first_paragraph


def _positioned_table_cells(page: IRPage) -> tuple[tuple[Any, IRBlock], ...]:
    cells: list[tuple[Any, IRBlock]] = []
    for block in page.body_blocks:
        if block.kind not in {"table", "html_table"}:
            continue
        table = getattr(block, "table", None)
        for cell in getattr(table, "cells", ()) or ():
            if str(getattr(cell, "text", "") or "").strip():
                cells.append((cell, block))
    return tuple(cells)


def _positioned_text_is_table_cell_duplicate(
    block: IRBlock,
    table_cells: tuple[tuple[Any, IRBlock], ...],
) -> bool:
    lines = tuple(block.lines or ())
    boxes = tuple(line.bbox for line in lines)
    if not boxes and block.bbox is not None:
        boxes = (block.bbox,)
    if not boxes:
        return False
    covered_line_count = 0
    for box in boxes:
        line_center = (
            (float(box[0]) + float(box[2])) / 2.0,
            (float(box[1]) + float(box[3])) / 2.0,
        )
        for cell, _ in table_cells:
            bbox = tuple(float(value) for value in cell.bbox)
            if (
                bbox[0] <= line_center[0] <= bbox[2]
                and bbox[1] <= line_center[1] <= bbox[3]
            ):
                covered_line_count += 1
                break
    return covered_line_count > 0


def _positioned_table_cell_lines(
    cell: Any,
) -> tuple[tuple[str, tuple[float, float, float, float], float], ...]:
    """按 glyph 的页面位置恢复密集表格单元格中的文字行。"""
    glyphs = tuple(getattr(cell, "glyphs", ()) or ())
    if glyphs:
        groups: list[list[Any]] = []
        centers: list[float] = []
        for glyph in sorted(glyphs, key=lambda item: (item.bbox[1], item.bbox[0])):
            center_y = (float(glyph.bbox[1]) + float(glyph.bbox[3])) / 2.0
            group_index = next(
                (
                    index
                    for index, value in enumerate(centers)
                    if abs(center_y - value) <= 3.0
                ),
                None,
            )
            if group_index is None:
                centers.append(center_y)
                groups.append([glyph])
            else:
                groups[group_index].append(glyph)
                centers[group_index] = sum(
                    (float(item.bbox[1]) + float(item.bbox[3])) / 2.0
                    for item in groups[group_index]
                ) / len(groups[group_index])
        result: list[tuple[str, tuple[float, float, float, float], float]] = []
        for group in sorted(groups, key=lambda items: min(item.bbox[1] for item in items)):
            ordered = sorted(group, key=lambda item: item.bbox[0])
            text = "".join(str(item.text or "") for item in ordered)
            bbox = (
                min(float(item.bbox[0]) for item in ordered),
                min(float(item.bbox[1]) for item in ordered),
                max(float(item.bbox[2]) for item in ordered),
                max(float(item.bbox[3]) for item in ordered),
            )
            size = max(float(getattr(item, "font_size", 0.0) or 0.0) for item in ordered)
            result.append((text, bbox, size))
        return tuple(result)
    raw_bbox = getattr(cell, "text_bbox", None) or getattr(cell, "bbox", None)
    if raw_bbox is None or len(raw_bbox) < 4:
        return ()
    return (
        (
            str(getattr(cell, "text", "") or "").strip(),
            tuple(float(value) for value in raw_bbox[:4]),
            float(getattr(cell, "font_size", 0.0) or 0.0),
        ),
    )


def _add_positioned_table_cell_text(
    document: Any,
    page: IRPage,
    table_block: IRBlock,
    *,
    cells: tuple[Any, ...],
    font_plan: FontPlan | None,
    canvas_paragraph: Any,
    report: dict[str, Any],
) -> None:
    fallback_font = str(getattr(table_block, "font_name", "") or "")
    for cell_index, cell in enumerate(cells):
        for line_index, (text, raw_bbox, glyph_size) in enumerate(
            _positioned_table_cell_lines(cell)
        ):
            text = text.strip()
            if not text:
                continue
            x0, top, width, height, fitted_size = _positioned_table_cell_metrics(
                cell,
                text,
                raw_bbox,
                glyph_size,
            )
            if width <= 0.0 or height <= 0.0:
                continue
            font_name, bold, italic = _planned_font(
                font_plan,
                raw_name="",
                fallback_name=fallback_font,
                bold=bool(getattr(cell, "bold", False)),
                italic=False,
            )
            add_absolute_text_box(
                document,
                x_points=x0,
                y_points=top,
                width_points=width,
                height_points=height,
                text=text,
                font_name=font_name,
                font_size=fitted_size,
                bold=bold,
                italic=italic,
                alignment=str(getattr(cell, "alignment", "") or "left"),
                object_id=6000 + len(report["placements"]) * 10000 + cell_index * 100 + line_index,
                name=f"positioned_table_cell_{table_block.block_id or 'table'}_{cell_index}_{line_index}",
                paragraph=canvas_paragraph,
            )


def _positioned_table_cell_metrics(
    cell: Any,
    text: str,
    raw_bbox: tuple[float, float, float, float],
    glyph_size: float,
) -> tuple[float, float, float, float, float]:
    """为单元格定位文本框保留字形行距，同时不越过单元格边界。"""
    x0, top, x1, bottom = (float(value) for value in raw_bbox)
    if x1 <= x0 or bottom <= top:
        return x0, top, 0.0, 0.0, 0.0
    base_size = float(glyph_size or getattr(cell, "font_size", 0.0) or 8.0)
    fitted_size = min(
        base_size,
        max(1.8, (x1 - x0) / max(len(text) * 0.68, 1.0)),
    )
    cell_bbox = tuple(
        float(value) for value in (getattr(cell, "bbox", None) or raw_bbox)
    )
    cell_right = max(cell_bbox[2], x0 + 1.0)
    width = min(max(x1 - x0 + 1.0, 1.0), max(cell_right - x0, 1.0))
    height = max(bottom - top + 1.0, fitted_size * 1.2, 1.0)
    return x0, top, width, height, fitted_size


def _add_positioned_form_block(
    document: Any,
    page: IRPage,
    block: IRBlock,
    *,
    source_pdf: Path | None,
    font_plan: FontPlan | None,
    canvas_paragraph: Any,
    max_fallback_pixels: int,
    report: dict[str, Any],
) -> Any | None:
    """渲染密集表单页，避免小块正文重新流式分页。"""
    if block.kind in FLOW_TEXT_KINDS:
        return _add_positioned_text_block(
            document,
            page,
            block,
            font_plan=font_plan,
            canvas_paragraph=canvas_paragraph,
            report=report,
        )
    if block.kind in {"table", "html_table"}:
        table = _add_table_block(
            document,
            page,
            block,
            source_pdf=source_pdf,
            margin_points=0.0,
            max_fallback_pixels=max_fallback_pixels,
            report=report,
            positioned=True,
        )
        if table is not None and block.bbox is not None:
            position_table(
                table,
                x_points=float(block.bbox[0]),
                y_points=float(block.bbox[1]),
                overlap=True,
            )
            _add_positioned_table_cell_text(
                document,
                page,
                block,
                cells=tuple(getattr(block.table, "cells", ()) or ()),
                font_plan=font_plan,
                canvas_paragraph=canvas_paragraph,
                report=report,
            )
            _record_placement(
                report,
                block,
                status="native",
                reason="positioned_table",
            )
        else:
            _record_placement(
                report,
                block,
                status="skipped",
                reason="positioned_table_unavailable",
            )
        return None
    if block.kind == "formula":
        return _add_positioned_text_block(
            document,
            page,
            block,
            font_plan=font_plan,
            canvas_paragraph=canvas_paragraph,
            report=report,
        )
    if block.kind in {"image", "page_image", "vector"}:
        image_bytes = block.image_bytes
        if image_bytes is None:
            image_bytes = _render_block_region(
                source_pdf,
                page,
                block,
                max_pixels=max_fallback_pixels,
            )
        if image_bytes is None or block.bbox is None:
            report["skipped_block_count"] += 1
            _record_placement(report, block, status="skipped", reason="positioned_image_unavailable")
            return None
        x0, top, x1, bottom = (float(value) for value in block.bbox)
        paragraph = canvas_paragraph
        add_absolute_picture(
            document,
            image_bytes,
            x_points=x0,
            y_points=top,
            width_points=max(x1 - x0, 1.0),
            height_points=max(bottom - top, 1.0),
            z_order=int(getattr(block, "z_order", 0) or 0),
            object_id=4000 + len(report["placements"]),
            name=f"positioned_{block.kind}_{block.block_id or 'block'}",
            paragraph=paragraph,
        )
        report["image_count"] += 1
        if block.kind == "page_image":
            report["page_image_count"] += 1
        _record_placement(report, block, status="native", reason="positioned_image")
        return paragraph
    report["skipped_block_count"] += 1
    _record_placement(report, block, status="skipped", reason="positioned_kind_unsupported")
    return None


def export_structured_docx(
    ir: IRDocument,
    output_path: Path,
    *,
    title: str | None = None,
    source_pdf: Path | None = None,
    include_toc: bool = False,
    include_bookmarks: bool = True,
    margin_points: float = FLOW_MARGIN_POINTS,
    max_fallback_pixels: int = 2_000_000,
    font_plan: FontPlan | None = None,
    stage_callback: StageCallback | None = None,
) -> dict[str, Any]:
    """按 IR 的逻辑块顺序生成可编辑、可流式重排的 DOCX。"""
    if margin_points < 0:
        raise ValueError("流式导出页边距不能为负数")
    if max_fallback_pixels < 1:
        raise ValueError("结构化导出图片像素上限必须大于 0")
    if source_pdf is not None:
        source_pdf = Path(source_pdf)

    document = Document()
    _set_document_styles(document)
    document.core_properties.title = title or ir.title
    _remove_initial_empty_paragraph(document)
    if include_toc:
        document.add_heading("目录", level=1)
        _add_toc_field(document)
        document.add_page_break()

    outline_by_page: dict[int, list[dict[str, Any]]] = {}
    if include_bookmarks:
        for entry in ir.metadata.get("outline", []):
            page_number = entry.get("page")
            if isinstance(page_number, int):
                outline_by_page.setdefault(page_number, []).append(entry)

    report: dict[str, Any] = {
        "mode": "structured",
        "page_count": len(ir.pages),
        "text_paragraph_count": 0,
        "text_image_fallback_count": 0,
        "rotated_text_image_fallback_count": 0,
        "table_count": 0,
        "html_table_count": 0,
        "formula_count": 0,
        "formula_omml_count": 0,
        "formula_image_fallback_count": 0,
        "formula_region_fallback_count": 0,
        "formula_text_fallback_count": 0,
        "image_count": 0,
        "page_image_count": 0,
        "page_image_fallback_count": 0,
        "visual_only_page_count": 0,
        "complex_table_region_fallback_count": 0,
        "vector_image_fallback_count": 0,
        "vector_compacted_page_count": 0,
        "vector_skipped_count": 0,
        "vector_background_shape_count": 0,
        "table_image_fallback_count": 0,
        "table_native_error_count": 0,
        "table_native_error_types": [],
        "header_footer_count": 0,
        "skipped_block_count": 0,
        "placements": [],
        "pages": [],
    }
    _notify_stage(stage_callback, "structured_export_started", page_count=len(ir.pages))
    bookmark_id = 0

    for page_index, page in enumerate(ir.pages):
        vector_blocks = [
            block for block in page.body_blocks if block.kind == "vector"
        ]
        dense_vector_bbox = (
            _vector_union_bbox(page)
            if len(vector_blocks) > DENSE_VECTOR_LIMIT
            else None
        )
        dense_vector_form_background = bool(
            dense_vector_bbox is not None
            and _dense_vector_is_form_background(page, dense_vector_bbox)
        )
        section = (
            document.sections[0]
            if page_index == 0
            else document.add_section(WD_SECTION.NEW_PAGE)
        )
        _set_section_page(
            section,
            max(page.width / 72.0, 0.01),
            max(page.height / 72.0, 0.01),
            0.0 if dense_vector_form_background else margin_points / 72.0,
        )
        page_image_fallback = _should_use_page_image_fallback(
            page,
            source_pdf=source_pdf,
        )
        if not page_image_fallback and (page.header_blocks or page.footer_blocks):
            _add_flow_header_footer(
                section,
                page,
                font_plan=font_plan,
                report=report,
            )
        else:
            detach_header_footer(section)

        start_index = len(report["placements"])
        if page_image_fallback:
            first_paragraph = _add_page_image_fallback(
                document,
                page,
                source_pdf=source_pdf,
                max_fallback_pixels=max_fallback_pixels,
                report=report,
            )
            if first_paragraph is not None:
                fallback_block = IRBlock(
                    kind="page_image",
                    page=page.page_number,
                    bbox=(0.0, 0.0, page.width, page.height),
                    source="complex_rotated_page_fallback",
                )
                _record_placement(
                    report,
                    fallback_block,
                    status="image_fallback",
                    reason="complex_rotated_page_rasterized",
                )
                if include_bookmarks:
                    for entry in outline_by_page.get(page.page_number, []):
                        bookmark_id += 1
                        _add_bookmark(
                            first_paragraph,
                            name=_bookmark_name(
                                str(entry.get("title") or ""),
                                bookmark_id,
                            ),
                            bookmark_id=bookmark_id,
                        )
                _record_page_report(report, page, start_index=start_index)
                _notify_stage(
                    stage_callback,
                    "structured_page_completed",
                    page=page.page_number,
                    page_count=len(ir.pages),
                )
                continue
        first_paragraph = None

        if dense_vector_form_background:
            positioned_canvas = new_canvas_paragraph(document)
            background_blocks = _dense_form_background_blocks(page)
            background_ids = {id(block) for block in background_blocks}
            for background_index, block in enumerate(background_blocks):
                vector = block.vector
                if vector is None or block.bbox is None:
                    continue
                x0, top, x1, bottom = (float(value) for value in block.bbox)
                add_absolute_shape(
                    document,
                    geometry="rect",
                    x_points=x0,
                    y_points=top,
                    width_points=max(x1 - x0, 1.0),
                    height_points=max(bottom - top, 1.0),
                    fill_color=getattr(vector, "fill_color", None),
                    line_color=None,
                    z_order=-100 + background_index,
                    behind_text=True,
                    object_id=3000 + background_index,
                    name=f"dense_form_background_{page.page_number}_{background_index}",
                    paragraph=positioned_canvas,
                )
                report["vector_background_shape_count"] += 1
                _record_placement(
                    report,
                    block,
                    status="native",
                    reason="dense_form_background_shape",
                )
            table_cells = _positioned_table_cells(page)
            report["positioned_table_cell_count"] = report.get(
                "positioned_table_cell_count", 0
            ) + len(table_cells)
            for block in page.body_blocks:
                if block.kind == "vector":
                    if id(block) in background_ids:
                        continue
                    report["vector_skipped_count"] += 1
                    report["skipped_block_count"] += 1
                    _record_placement(
                        report,
                        block,
                        status="skipped",
                        reason="dense_vector_form_background_skipped",
                    )
                    continue
                is_table_cell_duplicate = (
                    block.kind in FLOW_TEXT_KINDS
                    and _positioned_text_is_table_cell_duplicate(block, table_cells)
                )
                if is_table_cell_duplicate:
                    report["positioned_table_duplicate_block_count"] = report.get(
                        "positioned_table_duplicate_block_count", 0
                    ) + 1
                    _record_placement(
                        report,
                        block,
                        status="covered_by_table_cell",
                        reason="positioned_table_cell_text",
                    )
                    continue
                paragraph = _add_positioned_form_block(
                    document,
                    page,
                    block,
                    source_pdf=source_pdf,
                    font_plan=font_plan,
                    canvas_paragraph=positioned_canvas,
                    max_fallback_pixels=max_fallback_pixels,
                    report=report,
                )
                if first_paragraph is None and paragraph is not None:
                    first_paragraph = paragraph
            if first_paragraph is None:
                first_paragraph = positioned_canvas
            if include_bookmarks:
                for entry in outline_by_page.get(page.page_number, []):
                    bookmark_id += 1
                    _add_bookmark(
                        first_paragraph,
                        name=_bookmark_name(
                            str(entry.get("title") or ""),
                            bookmark_id,
                        ),
                        bookmark_id=bookmark_id,
                    )
            report["vector_compacted_page_count"] += 1
            _record_page_report(report, page, start_index=start_index)
            _notify_stage(
                stage_callback,
                "structured_page_completed",
                page=page.page_number,
                page_count=len(ir.pages),
            )
            continue

        dense_vector_image = None
        if dense_vector_bbox is not None:
            dense_vector_block = IRBlock(
                kind="vector",
                page=page.page_number,
                bbox=dense_vector_bbox,
                source="dense_vector_region",
            )
            dense_vector_image = _render_block_region(
                source_pdf,
                page,
                dense_vector_block,
                max_pixels=max_fallback_pixels,
            )
            report["vector_compacted_page_count"] += 1
        dense_vector_inserted = False
        complex_table_regions = []
        for region in _complex_table_regions(page):
            region_block = IRBlock(
                kind="image",
                page=page.page_number,
                bbox=region,
                image_width=region[2] - region[0],
                image_height=region[3] - region[1],
                image_alt="复杂表格区域",
                source="complex_table_region_fallback",
            )
            image_bytes = _render_block_region(
                source_pdf,
                page,
                region_block,
                max_pixels=max_fallback_pixels,
            )
            if image_bytes:
                region_block.image_bytes = image_bytes
                complex_table_regions.append(
                    {"region": region, "block": region_block, "inserted": False}
                )
        formula_regions = []
        for region in _formula_regions(page):
            region_block = IRBlock(
                kind="image",
                page=page.page_number,
                bbox=region,
                image_width=region[2] - region[0],
                image_height=region[3] - region[1],
                image_alt="复杂公式区域",
                source="formula_region_fallback",
            )
            image_bytes = _render_block_region(
                source_pdf,
                page,
                region_block,
                max_pixels=max_fallback_pixels,
            )
            if image_bytes:
                region_block.image_bytes = image_bytes
                formula_regions.append(
                    {
                        "region": region,
                        "block": region_block,
                        "inserted": False,
                        "formula_count": sum(
                            item.kind == "formula"
                            and _block_overlaps_region(item, region)
                            for item in page.body_blocks
                        ),
                    }
                )
        for block in page.body_blocks:
            table_region = next(
                (
                    item
                    for item in complex_table_regions
                    if _block_overlaps_region(block, item["region"])
                ),
                None,
            )
            if table_region is not None:
                if not table_region["inserted"]:
                    paragraph = _add_inline_image(
                        document,
                        table_region["block"],
                        page,
                        margin_points=margin_points,
                    )
                    if paragraph is not None:
                        table_region["inserted"] = True
                        report["table_image_fallback_count"] += 1
                        report["complex_table_region_fallback_count"] += 1
                        _record_placement(
                            report,
                            block,
                            status="image_fallback",
                            reason="complex_table_region_rasterized",
                        )
                        if first_paragraph is None:
                            first_paragraph = paragraph
                        continue
                elif table_region["inserted"]:
                    _record_placement(
                        report,
                        block,
                        status="covered_by_image_fallback",
                        reason="complex_table_region_rasterized",
                    )
                    continue
            formula_region = next(
                (
                    item
                    for item in formula_regions
                    if _block_overlaps_region(block, item["region"])
                ),
                None,
            )
            if formula_region is not None:
                if not formula_region["inserted"]:
                    paragraph = _add_inline_image(
                        document,
                        formula_region["block"],
                        page,
                        margin_points=margin_points,
                    )
                    if paragraph is not None:
                        formula_region["inserted"] = True
                        report["formula_count"] += formula_region["formula_count"]
                        report["formula_image_fallback_count"] += 1
                        report["formula_region_fallback_count"] += 1
                        _record_placement(
                            report,
                            block,
                            status="image_fallback",
                            reason="formula_region_rasterized",
                        )
                        if first_paragraph is None:
                            first_paragraph = paragraph
                        continue
                elif formula_region["inserted"]:
                    _record_placement(
                        report,
                        block,
                        status="covered_by_image_fallback",
                        reason="formula_region_rasterized",
                    )
                    continue
            if dense_vector_form_background and block.kind == "vector":
                report["vector_skipped_count"] += 1
                report["skipped_block_count"] += 1
                _record_placement(
                    report,
                    block,
                    status="skipped",
                    reason="dense_vector_form_background_skipped",
                )
                continue
            if dense_vector_bbox is not None and block.kind == "vector":
                if dense_vector_image is not None and not dense_vector_inserted:
                    image_block = IRBlock(
                        kind="image",
                        page=block.page,
                        bbox=dense_vector_bbox,
                        image_bytes=dense_vector_image,
                        image_width=dense_vector_bbox[2] - dense_vector_bbox[0],
                        image_height=dense_vector_bbox[3] - dense_vector_bbox[1],
                        image_alt="密集矢量区域",
                        source="dense_vector_fallback",
                    )
                    paragraph = _add_inline_image(
                        document,
                        image_block,
                        page,
                        margin_points=margin_points,
                    )
                    if paragraph is not None:
                        dense_vector_inserted = True
                        report["vector_image_fallback_count"] += 1
                        _record_placement(
                            report,
                            block,
                            status="image_fallback",
                            reason="dense_vector_layer_rasterized",
                        )
                        if first_paragraph is None:
                            first_paragraph = paragraph
                        continue
                report["vector_skipped_count"] += 1
                report["skipped_block_count"] += 1
                _record_placement(
                    report,
                    block,
                    status="skipped",
                    reason=(
                        "dense_vector_layer_compacted"
                        if dense_vector_image is not None
                        else "dense_vector_layer_unavailable"
                    ),
                )
                continue
            paragraph = _add_page_block(
                document,
                page,
                block,
                source_pdf=source_pdf,
                font_plan=font_plan,
                margin_points=margin_points,
                max_fallback_pixels=max_fallback_pixels,
                report=report,
            )
            if first_paragraph is None and paragraph is not None:
                first_paragraph = paragraph

        if first_paragraph is None:
            first_paragraph = document.add_paragraph()
            _configure_paragraph(first_paragraph)
        if include_bookmarks:
            for entry in outline_by_page.get(page.page_number, []):
                bookmark_id += 1
                _add_bookmark(
                    first_paragraph,
                    name=_bookmark_name(
                        str(entry.get("title") or ""),
                        bookmark_id,
                    ),
                    bookmark_id=bookmark_id,
                )
        _record_page_report(report, page, start_index=start_index)
        _notify_stage(
            stage_callback,
            "structured_page_completed",
            page=page.page_number,
            page_count=len(ir.pages),
        )

    if font_plan is not None and font_plan.programs:
        report["embedded_font_count"] = len(font_plan.programs)
        attach_embedded_fonts(document, font_plan.programs)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    _notify_stage(
        stage_callback,
        "structured_export_completed",
        page_count=len(ir.pages),
        text_paragraph_count=report["text_paragraph_count"],
        table_count=report["table_count"],
        formula_count=report["formula_count"],
        image_count=report["image_count"],
    )
    return report


__all__ = ["STRUCTURED_MODES", "export_structured_docx"]
