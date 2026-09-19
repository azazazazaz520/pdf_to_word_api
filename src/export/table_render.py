"""把版面提取出的 PdfTable 渲染为 Word 固定布局表格。

表格边框、列宽、行高与单元格内边距在此处对齐源 PDF；三条导出路径共用。
"""


from math import ceil
from typing import Any
from docx import Document
from docx.table import _Cell
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from ..layout.models import PdfTable


def _set_table_row_properties(row: Any, *, repeat_header: bool) -> None:
    row_properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    row_properties.append(cant_split)
    if repeat_header:
        table_header = OxmlElement("w:tblHeader")
        table_header.set(qn("w:val"), "true")
        row_properties.append(table_header)


def _set_table_cell_margins(
    table: Any,
    *,
    tight: bool = False,
    padding_points: tuple[float, float] | None = None,
) -> None:
    """把 Word 表格单元格上下左右边距压缩到接近源 PDF。"""
    table_properties = table._tbl.tblPr
    for existing in table_properties.findall(qn("w:tblCellMar")):
        table_properties.remove(existing)
    cell_margins = OxmlElement("w:tblCellMar")
    if padding_points is not None:
        left = max(int(round(float(padding_points[0]) * 20)), 0)
        top = max(int(round(float(padding_points[1]) * 20)), 0)
        values = (
            ("w:top", top),
            ("w:left", left),
            ("w:bottom", 0),
            ("w:right", left),
        )
    else:
        values = (
            (("w:top", 0), ("w:left", 10), ("w:bottom", 0), ("w:right", 10))
            if tight
            else (("w:top", 20), ("w:left", 40), ("w:bottom", 20), ("w:right", 40))
        )
    for tag, value in values:
        element = OxmlElement(tag)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        cell_margins.append(element)
    table_properties.append(cell_margins)


def _set_fixed_table_layout(table: Any) -> None:
    """将 Word 表格设置为固定布局，避免自动调整破坏 PDF 列宽比例。"""
    table_properties = table._tbl.tblPr
    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_cell_border(cell: Any, edge: str, value: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    element = borders.find(qn(f"w:{edge}"))
    if element is None:
        element = OxmlElement(f"w:{edge}")
        borders.append(element)
    element.set(qn("w:val"), value)
    element.set(qn("w:sz"), "0")


def _apply_visible_row_boundaries(table: Any, visibility: tuple[bool, ...]) -> None:
    if not visibility or len(visibility) != len(table.rows) + 1:
        return
    for boundary_index, is_visible in enumerate(visibility[1:-1], start=1):
        if is_visible:
            continue
        for cell in table.rows[boundary_index - 1].cells:
            _set_cell_border(cell, "bottom", "nil")
        for cell in table.rows[boundary_index].cells:
            _set_cell_border(cell, "top", "nil")


def _table_rows_needing_reflow(
    pdf_table: PdfTable,
    *,
    row_offset: int,
    cell_padding_points: tuple[float, float] | None,
) -> set[int]:
    """找出替代字体可能需要自动换行的表格行。"""
    cells = list(getattr(pdf_table, "cells", ()) or ())
    if not cells:
        return set()
    left_padding = (
        max(float(cell_padding_points[0]), 0.0)
        if cell_padding_points
        else 0.0
    )
    top_padding = (
        max(float(cell_padding_points[1]), 0.0)
        if cell_padding_points
        else 0.0
    )
    row_heights = pdf_table.row_heights
    reflow_rows: set[int] = set()
    for cell in cells:
        text = str(getattr(cell, "text", "") or "").strip()
        bbox = getattr(cell, "bbox", None)
        if not text or not bbox or len(bbox) < 4:
            continue
        font_size = max(float(getattr(cell, "font_size", 0.0) or 0.0), 0.0)
        if font_size <= 0:
            continue
        available_width = max(float(bbox[2]) - float(bbox[0]) - 2 * left_padding, 1.0)
        average_char_width = 0.6 * font_size
        estimated_width = sum(
            average_char_width * (0.5 if character.isspace() else 1.0)
            for character in text
        )
        required_lines = max(1, ceil(estimated_width / available_width))
        if required_lines <= 1:
            continue
        required_height = required_lines * max(font_size * 1.2, 6.0)
        cell_height = max(
            sum(
                row_heights[cell.row_index : cell.row_index + cell.row_span]
            )
            - top_padding,
            6.0,
        )
        if required_height <= cell_height + 1.0:
            continue
        start = cell.row_index + row_offset
        stop = start + max(int(cell.row_span), 1)
        reflow_rows.update(range(start, stop))
    return reflow_rows


def _set_pdf_cell_content(
    cell: Any,
    text: str,
    *,
    width: float,
    is_header: bool,
    font_size: float = 0.0,
    bold: bool = False,
    alignment: str = "",
    vertical_alignment: str = "center",
    spans: tuple[Any, ...] = (),
    shrink_cell_font: bool = False,
) -> None:
    """设置 PDF 表格单元格的宽度、对齐方式和文本格式。"""
    cell.width = Inches(width)
    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.TOP
        if vertical_alignment == "top"
        else WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )
    cell.text = ""
    alignments = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        if alignment in alignments:
            paragraph.alignment = alignments[alignment]
        if spans:
            previous_baseline: float | None = None
            for span in spans:
                span_glyphs = list(getattr(span, "glyphs", ()) or ())
                if span_glyphs:
                    groups: list[list[Any]] = []
                    baselines: list[float] = []
                    for glyph in span_glyphs:
                        baseline = float(glyph.bbox[3])
                        if (
                            not groups
                            or abs(baseline - baselines[-1]) <= 3.0
                        ):
                            if not groups:
                                groups.append([])
                                baselines.append(baseline)
                            groups[-1].append(glyph)
                            baselines[-1] = (
                                baselines[-1] * (len(groups[-1]) - 1)
                                + baseline
                            ) / len(groups[-1])
                        else:
                            groups.append([glyph])
                            baselines.append(baseline)
                    for group, baseline in zip(groups, baselines):
                        if (
                            previous_baseline is not None
                            and abs(baseline - previous_baseline) > 3.0
                        ):
                            paragraph.add_run().add_break()
                        run = paragraph.add_run(
                            "".join(str(glyph.text) for glyph in group)
                        )
                        span_font = str(getattr(span, "font_name", "") or "")
                        if span_font:
                            run.font.name = span_font
                        span_size = float(getattr(span, "font_size", 0.0) or 0.0)
                        group_width = max(
                            float(glyph.bbox[2]) for glyph in group
                        ) - min(float(glyph.bbox[0]) for glyph in group)
                        cell_width_points = width * 72.0
                        if (
                            shrink_cell_font
                            and span_size > 0
                            and cell_width_points > 0
                            and group_width >= cell_width_points * 0.78
                            and len(run.text.strip()) >= 6
                        ):
                            span_size = max(span_size * 0.9, 5.0)
                        if span_size > 0:
                            run.font.size = Pt(span_size)
                        run.bold = bool(getattr(span, "bold", False)) or is_header or bold
                        run.italic = bool(getattr(span, "italic", False))
                        color = getattr(span, "color", None)
                        if color and len(color) >= 3:
                            run.font.color.rgb = RGBColor(
                                int(color[0]) & 0xFF,
                                int(color[1]) & 0xFF,
                                int(color[2]) & 0xFF,
                            )
                        previous_baseline = baseline
                    continue
                run = paragraph.add_run(str(getattr(span, "text", "")))
                span_font = str(getattr(span, "font_name", "") or "")
                if span_font:
                    run.font.name = span_font
                span_size = float(getattr(span, "font_size", 0.0) or 0.0)
                if span_size > 0:
                    run.font.size = Pt(span_size)
                run.bold = bool(getattr(span, "bold", False)) or is_header or bold
                run.italic = bool(getattr(span, "italic", False))
                color = getattr(span, "color", None)
                if color and len(color) >= 3:
                    run.font.color.rgb = RGBColor(
                        int(color[0]) & 0xFF,
                        int(color[1]) & 0xFF,
                        int(color[2]) & 0xFF,
                    )
        else:
            run = paragraph.add_run(text)
            if is_header or bold:
                run.bold = True
            if font_size and font_size > 0:
                run.font.size = Pt(font_size)


def _cell_lookup(
    pdf_table: PdfTable,
    row_index: int,
    column_index: int,
) -> Any:
    for cell in pdf_table.cells:
        if (
            cell.row_index == row_index
            and cell.column_index == column_index
        ):
            return cell
    return None


def _cell_font_size(
    pdf_table: PdfTable,
    row_index: int,
    column_index: int,
) -> float:
    cell = _cell_lookup(pdf_table, row_index, column_index)
    return float(getattr(cell, "font_size", 0.0) or 0.0)


def _cell_alignment(
    pdf_table: PdfTable,
    row_index: int,
    column_index: int,
) -> str:
    cell = _cell_lookup(pdf_table, row_index, column_index)
    return str(getattr(cell, "alignment", "") or "")


def _physical_cell_grid(table: Any) -> list[list[_Cell]]:
    """按 Word XML 的物理行列保存单元格，避免纵向合并改变坐标映射。"""
    grid: list[list[_Cell]] = []
    for row in table.rows:
        row_cells: list[_Cell] = []
        grid_column = 0
        for tc in row._tr.tc_lst:
            tc_pr = tc.tcPr
            grid_span_element = (
                tc_pr.gridSpan
                if tc_pr is not None
                else None
            )
            grid_span = int(grid_span_element.val) if grid_span_element is not None else 1
            cell = _Cell(tc, row)
            row_cells.extend([cell] * max(grid_span, 1))
            grid_column += max(grid_span, 1)
        grid.append(row_cells)
    return grid


def _add_pdf_table(
    document: Document,
    pdf_table: PdfTable,
    *,
    bordered: bool = True,
    exact_widths: bool = False,
    cell_vertical_alignment: str = "center",
    tight_cell_margins: bool = False,
    cell_padding_points: tuple[float, float] | None = None,
    merge_cells: bool = True,
    empty_text: bool = False,
    exact_row_heights: bool = False,
) -> Any:
    if not pdf_table.rows or pdf_table.column_count < 1:
        return None
    table = document.add_table(rows=0, cols=pdf_table.column_count)
    if bordered:
        table.style = "Table Grid"
    table.autofit = False
    _set_fixed_table_layout(table)
    _set_table_cell_margins(
        table,
        tight=tight_cell_margins,
        padding_points=cell_padding_points,
    )
    widths = [
        max(
            (
                pdf_table.column_boundaries[index + 1]
                - pdf_table.column_boundaries[index]
            )
            / 72.0,
            0.02 if exact_widths else 0.2,
        )
        for index in range(pdf_table.column_count)
    ]
    if exact_widths:
        scale = 1.0
    else:
        section = document.sections[-1]
        available_width = max(
            section.page_width.inches
            - section.left_margin.inches
            - section.right_margin.inches,
            0.2,
        )
        scale = min(1.0, available_width / max(sum(widths), 0.2))
    widths = [width * scale for width in widths]
    continuation_header_rows = list(pdf_table.continuation_header_rows)
    source_rows = list(pdf_table.rows)
    row_offset = len(continuation_header_rows)
    all_rows = continuation_header_rows + source_rows
    header_row_count = row_offset + pdf_table.header_row_count
    row_heights = pdf_table.row_heights
    default_header_height = row_heights[0] if row_heights else 18.0
    all_row_heights = [
        *([default_header_height] * row_offset),
        *row_heights,
    ]
    # Word 会把单元格上内边距计入行高，这里换算成纯文本行高。
    top_padding = (
        max(float(cell_padding_points[1]), 0.0)
        if cell_padding_points
        else 0.0
    )
    all_row_heights = [
        max(height - top_padding, 6.0) for height in all_row_heights
    ]
    reflow_rows = (
        set()
        if exact_row_heights
        else _table_rows_needing_reflow(
            pdf_table,
            row_offset=row_offset,
            cell_padding_points=cell_padding_points,
        )
    )
    for row_index in range(len(all_rows)):
        row = table.add_row()
        _set_table_row_properties(
            row,
            repeat_header=row_index < header_row_count,
        )
        if row_index < len(all_row_heights):
            row.height = Pt(max(all_row_heights[row_index] * scale, 6.0))
            row.height_rule = (
                WD_ROW_HEIGHT_RULE.EXACTLY
                if exact_row_heights and row_index not in reflow_rows
                else WD_ROW_HEIGHT_RULE.AT_LEAST
            )
        for column_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for row_index, values in enumerate(continuation_header_rows):
        for column_index, cell in enumerate(table.rows[row_index].cells):
            _set_pdf_cell_content(
                cell,
                ""
                if empty_text
                else values[column_index] if column_index < len(values) else "",
                width=widths[column_index],
                is_header=True,
                font_size=_cell_font_size(pdf_table, row_index, column_index),
                alignment=_cell_alignment(pdf_table, row_index, column_index),
                vertical_alignment=cell_vertical_alignment,
                shrink_cell_font=exact_row_heights and row_index not in reflow_rows,
            )

    if pdf_table.cells:
        physical_cells = _physical_cell_grid(table)
        # 先完成网格合并，再按源坐标写入文字。逐个合并并立即写入会改变
        # python-docx 对后续行列的映射，导致纵向合并区域的文字落到上一行。
        if merge_cells:
            for pdf_cell in pdf_table.cells:
                if pdf_cell.row_span <= 1 and pdf_cell.column_span <= 1:
                    continue
                source_row_index = pdf_cell.row_index + row_offset
                try:
                    physical_cells[source_row_index][pdf_cell.column_index].merge(
                        physical_cells[
                            source_row_index + pdf_cell.row_span - 1
                        ][
                            pdf_cell.column_index + pdf_cell.column_span - 1
                        ]
                    )
                except Exception:
                    # 个别 PDF 的跨行跨列信息可能互相重叠；保留未合并网格，
                    # 让其余单元格仍能作为可编辑表格导出。
                    continue
        for pdf_cell in pdf_table.cells:
            source_row_index = pdf_cell.row_index + row_offset
            cell = physical_cells[source_row_index][pdf_cell.column_index]
            _set_pdf_cell_content(
                cell,
                "" if empty_text else pdf_cell.text,
                width=sum(
                    widths[
                        pdf_cell.column_index : pdf_cell.column_index
                        + pdf_cell.column_span
                    ]
                ),
                is_header=pdf_cell.row_index < pdf_table.header_row_count,
                font_size=pdf_cell.font_size,
                bold=pdf_cell.bold,
                alignment=pdf_cell.alignment,
                vertical_alignment=cell_vertical_alignment,
                spans=() if empty_text else pdf_cell.spans,
                shrink_cell_font=(
                    exact_row_heights and source_row_index not in reflow_rows
                ),
            )
    else:
        for row_index, values in enumerate(all_rows):
            for column_index, cell in enumerate(table.rows[row_index].cells):
                _set_pdf_cell_content(
                    cell,
                    ""
                    if empty_text
                    else values[column_index] if column_index < len(values) else "",
                    width=widths[column_index],
                    is_header=row_index < header_row_count,
                    font_size=_cell_font_size(pdf_table, row_index, column_index),
                    alignment=_cell_alignment(pdf_table, row_index, column_index),
                    vertical_alignment=cell_vertical_alignment,
                )
    _apply_visible_row_boundaries(
        table,
        tuple(getattr(pdf_table, "visible_row_boundaries", ()) or ()),
    )
    for column_index, column in enumerate(table.columns):
        column.width = Inches(widths[column_index])
    return table


