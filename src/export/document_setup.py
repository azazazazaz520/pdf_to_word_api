"""文档级设置与整页图片渲染：三条导出路径共用。

负责新建文档的样式、section 纸张与页边距、阶段回调，以及把源 PDF 页面
渲染成整页图片（供页面保真路径与 OCR 页使用）。
"""

from __future__ import annotations

import math
from io import BytesIO
from pathlib import Path
from collections.abc import Callable
from typing import Any
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image


DEFAULT_PAGE_IMAGE_MAX_PIXELS = 4 * 1024 * 1024
DEFAULT_PAGE_IMAGE_JPEG_QUALITY = 88
PAGE_IMAGE_MAX_DIMENSION_INCHES = 11.0
"""页面尺寸换算时允许的最大边长（英寸）；超出后等比缩小。"""
StageCallback = Callable[[str, dict[str, Any]], None]


def _set_document_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(9.5)
    for style_name in ("Normal", "List Paragraph", "List Bullet", "List Number"):
        style = document.styles[style_name]
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.bold = True
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1
    if "Code Block" not in [style.name for style in document.styles]:
        code_style = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(8.0)
        code_style.paragraph_format.space_before = Pt(0)
        code_style.paragraph_format.space_after = Pt(0)
        code_style.paragraph_format.line_spacing = 1
        code_rpr = code_style.element.get_or_add_rPr()
        code_fonts = code_rpr.get_or_add_rFonts()
        code_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _notify_stage(
    callback: StageCallback | None, stage: str, **details: Any
) -> None:
    if callback is not None:
        callback(stage, details)


def _set_section_page(
    section: Any,
    width_inches: float,
    height_inches: float,
    margin_inches: float,
) -> None:
    from docx.shared import Inches

    section.page_width = Inches(width_inches)
    section.page_height = Inches(height_inches)
    section.top_margin = Inches(margin_inches)
    section.bottom_margin = Inches(margin_inches)
    section.left_margin = Inches(margin_inches)
    section.right_margin = Inches(margin_inches)


def _scaled_page_size(
    width_points: float, height_points: float
) -> tuple[float, float]:
    width_inches = max(width_points / 72.0, 0.01)
    height_inches = max(height_points / 72.0, 0.01)
    scale = min(1.0, PAGE_IMAGE_MAX_DIMENSION_INCHES / max(width_inches, height_inches))
    return width_inches * scale, height_inches * scale


def _encode_page_image(
    image: Image.Image,
    *,
    max_pixels: int,
    jpeg_quality: int,
) -> bytes:
    image = image.convert("RGB")
    pixel_count = image.width * image.height
    if pixel_count > max_pixels:
        scale = math.sqrt(max_pixels / pixel_count)
        target_width = max(1, int(image.width * scale))
        target_height = max(1, int(image.height * scale))
        while target_width * target_height > max_pixels:
            if target_width >= target_height:
                target_width -= 1
            else:
                target_height -= 1
        image = image.resize(
            (max(1, target_width), max(1, target_height)),
            Image.Resampling.LANCZOS,
        )
    output = BytesIO()
    image.save(
        output,
        format="JPEG",
        quality=max(1, min(100, jpeg_quality)),
        optimize=False,
    )
    return output.getvalue()


def render_page_image_png(
    source_pdf: Path,
    page_index: int,
    *,
    dpi: float = 110.0,
    max_pixels: int = 8_000_000,
) -> bytes:
    """按指定 DPI 渲染整页 PNG，用于高保真自动兜底页。"""
    import pypdfium2 as pdfium

    if not source_pdf.is_file():
        raise FileNotFoundError(f"页面渲染所需的 PDF 不存在：{source_pdf}")
    scale = max(0.2, min(float(dpi) / 72.0, 6.0))
    document = pdfium.PdfDocument(str(source_pdf))
    page = document[page_index]
    try:
        bitmap = page.render(scale=scale)
        try:
            image = bitmap.to_pil()
            try:
                pixel_count = image.width * image.height
                if pixel_count > max_pixels > 0:
                    ratio = math.sqrt(max_pixels / pixel_count)
                    image = image.resize(
                        (
                            max(1, int(image.width * ratio)),
                            max(1, int(image.height * ratio)),
                        ),
                        Image.Resampling.LANCZOS,
                    )
                output = BytesIO()
                image.save(output, format="PNG", optimize=False)
                return output.getvalue()
            finally:
                image.close()
        finally:
            close_bitmap = getattr(bitmap, "close", None)
            if close_bitmap is not None:
                close_bitmap()
    finally:
        page.close()
        document.close()


def page_render_scale(
    width_points: float,
    height_points: float,
    max_pixels: int,
) -> float:
    """页面保真渲染使用的缩放系数，供 OCR bbox 换算复用。"""
    area = max(float(width_points) * float(height_points), 1.0)
    return max(min(2.0, math.sqrt(max(int(max_pixels), 1) / area)), 0.1)


def _render_page_image(
    pdf_document: Any,
    page_index: int,
    *,
    max_pixels: int,
    jpeg_quality: int,
) -> bytes:
    import pypdfium2 as pdfium

    pdf_page = pdf_document[page_index]
    try:
        width, height = pdf_page.get_size()
        scale = page_render_scale(width, height, max_pixels)
        bitmap = pdf_page.render(scale=scale)
        try:
            image = bitmap.to_pil()
            return _encode_page_image(
                image,
                max_pixels=max_pixels,
                jpeg_quality=jpeg_quality,
            )
        finally:
            close_bitmap = getattr(bitmap, "close", None)
            if close_bitmap is not None:
                close_bitmap()
    finally:
        pdf_page.close()


def _remove_initial_empty_paragraph(document: Document) -> None:
    if len(document.paragraphs) != 1 or document.paragraphs[0].text:
        return
    paragraph = document.paragraphs[0]._element
    paragraph.getparent().remove(paragraph)


def render_pdf_region(
    source_pdf: Path,
    page_index: int,
    bbox: tuple[float, float, float, float],
    *,
    max_pixels: int = 1_500_000,
    jpeg_quality: int = 90,
    scale: float | None = None,
) -> bytes:
    """渲染 PDF 页面指定区域，用于公式图片和低置信度区域兜底。"""
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(source_pdf))
    page = document[page_index]
    try:
        render_scale = float(scale) if scale else 3.0
        render_scale = max(0.5, min(render_scale, 6.0))
        bitmap = page.render(scale=render_scale)
        try:
            image = bitmap.to_pil()
            x0, top, x1, bottom = bbox
            left = max(0, int(x0 * render_scale))
            upper = max(0, int(top * render_scale))
            right = min(image.width, int(x1 * render_scale))
            lower = min(image.height, int(bottom * render_scale))
            if right <= left or lower <= upper:
                raise ValueError("公式区域为空")
            crop = image.crop((left, upper, right, lower))
            if crop.width < 8 or crop.height < 8:
                raise ValueError("公式区域过小")
            if crop.width * crop.height > max_pixels:
                scale = math.sqrt(max_pixels / (crop.width * crop.height))
                crop = crop.resize(
                    (
                        max(1, int(crop.width * scale)),
                        max(1, int(crop.height * scale)),
                    ),
                    Image.Resampling.LANCZOS,
                )
            output = BytesIO()
            crop.convert("RGB").save(
                output,
                format="JPEG",
                quality=max(1, min(100, jpeg_quality)),
                optimize=False,
            )
            return output.getvalue()
        finally:
            close_bitmap = getattr(bitmap, "close", None)
            if close_bitmap is not None:
                close_bitmap()
    finally:
        page.close()
        document.close()
