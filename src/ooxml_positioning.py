"""OOXML 绝对定位原语。

高保真导出把每一页当作一块画布：所有内容都通过 page-relative 的
绝对坐标写入，不再参与 Word 的流式排版。本模块集中封装所需的
OOXML/DrawingML 结构：

* 文本：``w:framePr`` 绝对定位；复杂场景使用 ``wps`` 文本框
* 图片：``wp:anchor`` 页面锚定
* 线条/矢量图形：``wps`` 形状，VML 兜底
* 表格：``w:tblpPr`` + 固定布局
* 页眉页脚：写入 header/footer part，定位方式与正文一致

导出方（``pdf_to_word_exporter``、``export/fidelity.py``）使用的公开原语：

* 尺寸与颜色：``set_exact_page``、``points_to_emu``、``points_to_twips``、``rgb_to_hex``
* 文本：``add_absolute_text_paragraph``、``add_absolute_text_box``、``style_run``
* 图形：``add_absolute_picture``、``add_absolute_shape``
* 表格：``position_table``
* 页眉页脚：``prepare_header``、``prepare_footer``、``header_footer_paragraph``、
  ``detach_header_footer``
* 页面承载：``new_canvas_paragraph``、``make_flow_paragraph_minimal``

其余以下划线开头的函数与常量是本模块内部工具（如 ``_frame_element``、
``_anchor_drawing``、``_shape_xml``、``_PPR_ORDER``），不构成对外接口；
判断某个函数能否删除时，必须同时统计模块内调用，不能只看模块外引用。
"""

from __future__ import annotations

import math
from copy import deepcopy
from io import BytesIO
from typing import Any, Iterable, Sequence

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt


EMU_PER_POINT = 12700
TWIPS_PER_POINT = 20
WORD_MAX_PAGE_INCHES = 22.0
"""Word 允许的最大页面尺寸；超过时需要等比缩小并记录告警。"""

_ALIGNMENT_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "start": WD_ALIGN_PARAGRAPH.LEFT,
    "end": WD_ALIGN_PARAGRAPH.RIGHT,
}

_WML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_WPS_NS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
_V_NS = "urn:schemas-microsoft-com:vml"
_W10_NS = "urn:schemas-microsoft-com:office:word"

_PPR_ORDER = (
    "w:pStyle",
    "w:keepNext",
    "w:keepLines",
    "w:pageBreakBefore",
    "w:framePr",
    "w:widowControl",
    "w:numPr",
    "w:suppressLineNumbers",
    "w:pBdr",
    "w:shd",
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:kinsoku",
    "w:wordWrap",
    "w:overflowPunct",
    "w:topLinePunct",
    "w:autoSpaceDE",
    "w:autoSpaceDN",
    "w:bidi",
    "w:adjustRightInd",
    "w:snapToGrid",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
)

_RPR_ORDER = (
    "w:rStyle",
    "w:rFonts",
    "w:b",
    "w:bCs",
    "w:i",
    "w:iCs",
    "w:caps",
    "w:smallCaps",
    "w:strike",
    "w:dstrike",
    "w:outline",
    "w:shadow",
    "w:emboss",
    "w:imprint",
    "w:noProof",
    "w:snapToGrid",
    "w:vanish",
    "w:webHidden",
    "w:color",
    "w:spacing",
    "w:w",
    "w:kern",
    "w:position",
    "w:sz",
    "w:szCs",
    "w:highlight",
    "w:u",
    "w:effect",
    "w:bdr",
    "w:shd",
    "w:fitText",
    "w:vertAlign",
    "w:rtl",
    "w:cs",
    "w:em",
    "w:lang",
    "w:eastAsianLayout",
    "w:specVanish",
    "w:oMath",
)


def points_to_emu(points: float) -> int:
    return int(round(float(points) * EMU_PER_POINT))


def points_to_twips(points: float) -> int:
    return int(round(float(points) * TWIPS_PER_POINT))


def rgb_to_hex(color: Any, default: str = "000000") -> str:
    """把 (r, g, b) 元组或 "RRGGBB" 字符串转成 Word 颜色值。"""
    if color is None:
        return default
    if isinstance(color, str):
        value = color.strip().lstrip("#")
        if len(value) == 6 and all(
            character in "0123456789abcdefABCDEF" for character in value
        ):
            return value.upper()
        return default
    try:
        red, green, blue = (int(max(0, min(255, value))) for value in color[:3])
    except Exception:
        return default
    return f"{red:02X}{green:02X}{blue:02X}"


def insert_ordered(parent: Any, element: Any, order: Sequence[str]) -> Any:
    """按 OOXML schema 顺序把元素插入到父元素中。

    顺序错误的 ``w:framePr`` / ``w:tblpPr`` 会被 Word 直接忽略，
    因此这里统一处理。
    """

    tag = element.tag
    index = len(parent)
    for position, child in enumerate(parent):
        child_tag = child.tag
        if child_tag == tag:
            parent.remove(child)
            index = position
            break
        if child_tag in order and order.index(child_tag) > order.index(tag):
            index = position
            break
    parent.insert(index, element)
    return element


def _set_ordered_ppr_child(paragraph: Any, element: Any) -> Any:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    return insert_ordered(paragraph_properties, element, _PPR_ORDER)


def set_exact_page(
    section: Any,
    *,
    width_points: float,
    height_points: float,
    margin_points: float = 0.0,
    header_distance_points: float = 0.0,
    footer_distance_points: float = 0.0,
) -> dict[str, Any]:
    """把 section 的纸张和页边距设置为源 PDF 的精确尺寸。

    返回 ``{"width_points", "height_points", "scaled"}``；仅当页面
    超过 Word 上限时才等比缩小。
    """

    width = max(float(width_points), 1.0)
    height = max(float(height_points), 1.0)
    scaled = False
    limit = WORD_MAX_PAGE_INCHES * 72.0
    largest = max(width, height)
    if largest > limit:
        ratio = limit / largest
        width *= ratio
        height *= ratio
        scaled = True
    section.page_width = Pt(width)
    section.page_height = Pt(height)
    section.top_margin = Pt(max(margin_points, 0.0))
    section.bottom_margin = Pt(max(margin_points, 0.0))
    section.left_margin = Pt(max(margin_points, 0.0))
    section.right_margin = Pt(max(margin_points, 0.0))
    section.header_distance = Pt(max(header_distance_points, 0.0))
    section.footer_distance = Pt(max(footer_distance_points, 0.0))
    return {"width_points": width, "height_points": height, "scaled": scaled}


def make_flow_paragraph_minimal(paragraph: Any, *, font_size: float = 1.0) -> Any:
    """把承载浮动对象的段落压缩到最小高度，避免流式排版产生额外页面。"""
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing = Pt(max(font_size, 0.5))
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    for run in paragraph.runs:
        run.font.size = Pt(font_size)
    paragraph_mark = paragraph._p.get_or_add_pPr()
    mark_properties = paragraph_mark.find(qn("w:rPr"))
    if mark_properties is None:
        mark_properties = OxmlElement("w:rPr")
        insert_ordered(paragraph_mark, mark_properties, _PPR_ORDER)
    size = mark_properties.find(qn("w:sz"))
    if size is None:
        size = OxmlElement("w:sz")
        insert_ordered(mark_properties, size, _RPR_ORDER)
    size.set(qn("w:val"), str(int(round(max(font_size, 0.5) * 2))))
    return paragraph


def clear_paragraph(paragraph: Any) -> Any:
    """删除段落中的所有 run，保留段落本身。"""
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    return paragraph


def new_canvas_paragraph(container: Any) -> Any:
    """新建一个只承载浮动对象的零高度画布段落。"""
    paragraph = container.add_paragraph()
    return make_flow_paragraph_minimal(paragraph)


def _frame_element(
    *,
    x_points: float,
    y_points: float,
    width_points: float | None,
    height_points: float | None = None,
    wrap: str = "none",
    h_anchor: str = "page",
    v_anchor: str = "page",
) -> Any:
    frame = OxmlElement("w:framePr")
    frame.set(qn("w:wrap"), wrap)
    frame.set(qn("w:hAnchor"), h_anchor)
    frame.set(qn("w:vAnchor"), v_anchor)
    frame.set(qn("w:x"), str(points_to_twips(x_points)))
    frame.set(qn("w:y"), str(points_to_twips(y_points)))
    if width_points is not None:
        frame.set(qn("w:w"), str(max(points_to_twips(width_points), 20)))
    if height_points is not None:
        frame.set(qn("w:h"), str(max(points_to_twips(height_points), 20)))
        frame.set(qn("w:hRule"), "atLeast")
    else:
        frame.set(qn("w:hRule"), "auto")
    return frame


def apply_absolute_frame(
    paragraph: Any,
    *,
    x_points: float,
    y_points: float,
    width_points: float | None,
    height_points: float | None = None,
    wrap: str = "none",
) -> Any:
    """给段落加上 ``w:framePr``，实现相对页面的绝对定位。"""
    frame = _frame_element(
        x_points=x_points,
        y_points=y_points,
        width_points=width_points,
        height_points=height_points,
        wrap=wrap,
    )
    _set_ordered_ppr_child(paragraph, frame)
    return paragraph


def set_run_size_half_points(run: Any, points: float) -> Any:
    """按四舍五入写入 ``w:sz``/``w:szCs``。

    python-docx 的 ``run.font.size = Pt(x)`` 会把半磅值**向下取整**，
    例如 10.45pt 会写成 20 半磅（10.0pt），导致整行窄 4%+；
    这里覆盖成 round(x*2)，保证字号误差 < 0.25pt。
    """
    value = max(int(round(float(points) * 2.0)), 1)
    run_properties = run._element.get_or_add_rPr()
    for tag in ("w:sz", "w:szCs"):
        element = run_properties.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            insert_ordered(run_properties, element, _RPR_ORDER)
        element.set(qn("w:val"), str(value))
    return run


def style_run(
    run: Any,
    *,
    font_name: str = "",
    font_size: float = 0.0,
    bold: bool = False,
    italic: bool = False,
    color: Any = None,
    underline: bool = False,
) -> Any:
    """设置 run 的字体信息，并同步 eastAsia 字体，避免中文回退。"""
    if font_name:
        run.font.name = font_name
        run_properties = run._element.get_or_add_rPr()
        fonts = run_properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            insert_ordered(run_properties, fonts, _RPR_ORDER)
        for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attribute), font_name)
    if font_size and font_size > 0:
        run.font.size = Pt(font_size)
        set_run_size_half_points(run, font_size)
    run.font.bold = bool(bold)
    run.font.italic = bool(italic)
    if underline:
        run.font.underline = True
    if color is not None:
        from docx.shared import RGBColor

        hex_value = rgb_to_hex(color)
        run.font.color.rgb = RGBColor.from_string(hex_value)
    return run


def add_absolute_text_paragraph(
    container: Any,
    text: str,
    *,
    x_points: float,
    y_points: float,
    width_points: float | None,
    height_points: float | None = None,
    font_name: str = "",
    font_size: float = 0.0,
    bold: bool = False,
    italic: bool = False,
    color: Any = None,
    alignment: str = "left",
    line_spacing_points: float | None = None,
    first_line_indent_points: float = 0.0,
    space_before_points: float = 0.0,
    space_after_points: float = 0.0,
    character_spacing_points: float = 0.0,
    paragraph: Any = None,
) -> Any:
    """写入一段绝对定位文本；文本通过 ``w:framePr`` 固定在页面上。"""
    target = paragraph if paragraph is not None else container.add_paragraph()
    paragraph_format = target.paragraph_format
    paragraph_format.space_before = Pt(max(space_before_points, 0.0))
    paragraph_format.space_after = Pt(max(space_after_points, 0.0))
    size = float(font_size) if font_size and font_size > 0 else 10.0
    spacing = float(line_spacing_points) if line_spacing_points else size
    paragraph_format.line_spacing = Pt(max(spacing, 0.5))
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    target.alignment = _ALIGNMENT_MAP.get(str(alignment).lower(), WD_ALIGN_PARAGRAPH.LEFT)
    # 关闭"中文与西文/数字之间自动加空格"，否则与 PDF 原坐标不一致
    ppr = target._p.get_or_add_pPr()
    for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):
        element = ppr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            insert_ordered(ppr, element, _PPR_ORDER)
        element.set(qn("w:val"), "0")
    if first_line_indent_points and first_line_indent_points > 0:
        paragraph_format.first_line_indent = Pt(first_line_indent_points)
    apply_absolute_frame(
        target,
        x_points=x_points,
        y_points=y_points,
        width_points=width_points,
        height_points=height_points,
    )
    if text:
        parts = text.split("\n")
        for index, part in enumerate(parts):
            if index:
                target.add_run().add_break()
            run = target.add_run(part)
            style_run(
                run,
                font_name=font_name,
                font_size=size,
                bold=bold,
                italic=italic,
                color=color,
            )
            if character_spacing_points:
                apply_character_spacing(run, character_spacing_points)
    return target


def apply_character_spacing(run: Any, points: float) -> Any:
    """设置 ``w:spacing``（字符间距，1/20 pt）。"""
    value = int(round(float(points) * 20))
    if value == 0:
        return run
    run_properties = run._element.get_or_add_rPr()
    for existing in run_properties.findall(qn("w:spacing")):
        run_properties.remove(existing)
    element = OxmlElement("w:spacing")
    element.set(qn("w:val"), str(value))
    insert_ordered(run_properties, element, _RPR_ORDER)
    return run


def _anchor_drawing(
    run: Any,
    *,
    x_points: float,
    y_points: float,
    width_points: float,
    height_points: float,
    z_order: int,
    object_id: int,
    name: str,
    behind_text: bool,
    rotation: float = 0.0,
) -> None:
    """把 ``run.add_picture`` 生成的 inline drawing 转成页面锚定 drawing。"""
    drawing = run._element.find(qn("w:drawing"))
    if drawing is None:
        return
    inline = drawing[0]
    anchor = OxmlElement("wp:anchor")
    anchor.set("distT", "0")
    anchor.set("distB", "0")
    anchor.set("distL", "0")
    anchor.set("distR", "0")
    anchor.set("simplePos", "0")
    anchor.set("relativeHeight", str(max(1, 251658240 + int(z_order) * 100)))
    anchor.set("behindDoc", "1" if behind_text else "0")
    anchor.set("locked", "0")
    anchor.set("layoutInCell", "1")
    anchor.set("allowOverlap", "1")
    simple_position = OxmlElement("wp:simplePos")
    simple_position.set("x", "0")
    simple_position.set("y", "0")
    anchor.append(simple_position)
    horizontal = OxmlElement("wp:positionH")
    horizontal.set("relativeFrom", "page")
    horizontal_offset = OxmlElement("wp:posOffset")
    horizontal_offset.text = str(points_to_emu(x_points))
    horizontal.append(horizontal_offset)
    anchor.append(horizontal)
    vertical = OxmlElement("wp:positionV")
    vertical.set("relativeFrom", "page")
    vertical_offset = OxmlElement("wp:posOffset")
    vertical_offset.text = str(points_to_emu(y_points))
    vertical.append(vertical_offset)
    anchor.append(vertical)
    for tag in ("wp:extent", "wp:effectExtent"):
        element = inline.find(qn(tag))
        if element is not None:
            anchor.append(element)
    anchor.append(OxmlElement("wp:wrapNone"))
    document_properties = inline.find(qn("wp:docPr"))
    if document_properties is not None:
        document_properties.set("id", str(max(1, int(object_id))))
        document_properties.set("name", name or "Picture")
        anchor.append(document_properties)
    frame_properties = inline.find(qn("wp:cNvGraphicFramePr"))
    if frame_properties is not None:
        anchor.append(frame_properties)
    graphic = inline.find(qn("a:graphic"))
    if graphic is not None:
        anchor.append(graphic)
    if rotation:
        _apply_picture_rotation(anchor, rotation)
    drawing.replace(inline, anchor)


def _apply_picture_rotation(anchor: Any, rotation: float) -> None:
    transform = anchor.find(".//" + qn("a:xfrm"))
    if transform is not None:
        transform.set("rot", str(int(round(float(rotation) * 60000))))


def add_absolute_picture(
    container: Any,
    image_bytes: bytes,
    *,
    x_points: float,
    y_points: float,
    width_points: float,
    height_points: float,
    z_order: int = 0,
    object_id: int = 1,
    name: str = "",
    behind_text: bool = False,
    rotation: float = 0.0,
    paragraph: Any = None,
    run: Any = None,
) -> Any:
    """按页面绝对坐标插入图片。"""
    target_paragraph = paragraph
    if target_paragraph is None:
        target_paragraph = container.add_paragraph()
        make_flow_paragraph_minimal(target_paragraph)
    target_run = run if run is not None else target_paragraph.add_run()
    target_run.add_picture(
        BytesIO(image_bytes),
        width=Pt(max(width_points, 0.5)),
        height=Pt(max(height_points, 0.5)),
    )
    _anchor_drawing(
        target_run,
        x_points=x_points,
        y_points=y_points,
        width_points=width_points,
        height_points=height_points,
        z_order=z_order,
        object_id=object_id,
        name=name,
        behind_text=behind_text,
        rotation=rotation,
    )
    return target_run


def _shape_xml(
    *,
    object_id: int,
    name: str,
    x_points: float,
    y_points: float,
    width_points: float,
    height_points: float,
    z_order: int,
    behind_text: bool,
    rotation: float,
    geometry: str,
    fill_hex: str | None,
    line_hex: str | None,
    line_width_points: float,
    flip_h: bool = False,
    flip_v: bool = False,
) -> bytes:
    width_emu = max(points_to_emu(width_points), 1)
    height_emu = max(points_to_emu(height_points), 1)
    extent_emu = max(points_to_emu(max(width_points, height_points)), 1)
    fill = (
        f'<a:solidFill><a:srgbClr val="{fill_hex}"/></a:solidFill>'
        if fill_hex
        else "<a:noFill/>"
    )
    line = (
        f'<a:ln w="{max(points_to_emu(line_width_points), 0)}">'
        f'<a:solidFill><a:srgbClr val="{line_hex}"/></a:solidFill></a:ln>'
        if line_hex
        else '<a:ln><a:noFill/></a:ln>'
    )
    rotation_attr = f' rot="{int(round(rotation * 60000))}"' if rotation else ""
    if flip_h:
        rotation_attr += ' flipH="1"'
    if flip_v:
        rotation_attr += ' flipV="1"'
    vml_style = (
        f"position:absolute;left:{x_points}pt;top:{y_points}pt;"
        f"width:{width_points}pt;height:{height_points}pt;"
        f"mso-position-horizontal-relative:page;"
        f"mso-position-vertical-relative:page;z-index:{z_order}"
    )
    vml_geometry = (
        f'<v:line from="0,0" to="{width_points}pt,{height_points}pt" '
        f'style="{vml_style}" strokecolor="#{line_hex or "000000"}" '
        f'strokeweight="{line_width_points}pt"/>'
        if geometry == "line"
        else f'<v:rect style="{vml_style}" fillcolor="#{fill_hex or "FFFFFF"}" stroked="{"t" if line_hex else "f"}" strokecolor="#{line_hex or "000000"}"/>'
    )
    return (
        f'<w:r xmlns:w="{_WML_NS}" xmlns:wp="{_WP_NS}" xmlns:a="{_A_NS}" '
        f'xmlns:mc="{_MC_NS}" xmlns:wps="{_WPS_NS}" xmlns:v="{_V_NS}" '
        f'xmlns:w10="{_W10_NS}">'
        f"<mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing>"
        f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="{max(1, 251658240 + int(z_order) * 100)}" '
        f'behindDoc="{"1" if behind_text else "0"}" locked="0" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{points_to_emu(x_points)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{points_to_emu(y_points)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{width_emu}" cy="{height_emu}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f"<wp:wrapNone/>"
        f'<wp:docPr id="{max(1, int(object_id))}" name="{name or "Shape"}"/>'
        f"<wp:cNvGraphicFramePr/>"
        f"<a:graphic><a:graphicData uri=\"{_WPS_NS}\">"
        f"<wps:wsp><wps:cNvSpPr/><wps:spPr>"
        f'<a:xfrm{rotation_attr}><a:off x="0" y="0"/><a:ext cx="{width_emu}" cy="{height_emu}"/></a:xfrm>'
        f'<a:prstGeom prst="{geometry}"><a:avLst/></a:prstGeom>'
        f"{fill}{line}"
        f"</wps:spPr><wps:bodyPr/></wps:wsp>"
        f"</a:graphicData></a:graphic>"
        f"</wp:anchor></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict>{vml_geometry}</w:pict></mc:Fallback>"
        f"</mc:AlternateContent></w:r>"
    ).encode("utf-8")


def add_absolute_shape(
    container: Any,
    *,
    geometry: str,
    x_points: float,
    y_points: float,
    width_points: float,
    height_points: float,
    fill_color: Any = None,
    line_color: Any = None,
    line_width_points: float = 0.75,
    z_order: int = 0,
    behind_text: bool = False,
    rotation: float = 0.0,
    object_id: int = 1,
    name: str = "",
    paragraph: Any = None,
    flip_h: bool = False,
    flip_v: bool = False,
) -> Any:
    """插入一个绝对定位的矢量形状（rect / line）。"""
    target_paragraph = paragraph
    if target_paragraph is None:
        target_paragraph = container.add_paragraph()
        make_flow_paragraph_minimal(target_paragraph)
    run = parse_xml(
        _shape_xml(
            object_id=object_id,
            name=name,
            x_points=x_points,
            y_points=y_points,
            width_points=width_points,
            height_points=height_points,
            z_order=z_order,
            behind_text=behind_text,
            rotation=rotation,
            geometry=geometry if geometry in {"rect", "line"} else "rect",
            fill_hex=rgb_to_hex(fill_color) if fill_color is not None else None,
            line_hex=rgb_to_hex(line_color) if line_color is not None else None,
            line_width_points=line_width_points,
            flip_h=flip_h,
            flip_v=flip_v,
        )
    )
    target_paragraph._p.append(run)
    return target_paragraph


def add_absolute_text_box(
    container: Any,
    *,
    x_points: float,
    y_points: float,
    width_points: float,
    height_points: float,
    text: str,
    font_name: str = "",
    font_size: float = 10.0,
    bold: bool = False,
    italic: bool = False,
    color: Any = None,
    alignment: str = "left",
    rotation: float = 0.0,
    z_order: int = 0,
    object_id: int = 1,
    name: str = "",
    paragraph: Any = None,
) -> Any:
    """复杂场景（旋转、无法用 framePr 描述）下的文本框实现。"""
    target_paragraph = paragraph
    if target_paragraph is None:
        target_paragraph = container.add_paragraph()
        make_flow_paragraph_minimal(target_paragraph)
    width_emu = max(points_to_emu(width_points), 1)
    height_emu = max(points_to_emu(height_points), 1)
    color_hex = rgb_to_hex(color) if color is not None else "000000"
    box_name = name or "TextBox"
    alignment_value = {
        "center": "center",
        "right": "right",
        "justify": "both",
    }.get(str(alignment).lower(), "left")
    font_size_half_points = max(int(round(float(font_size) * 2)), 1)
    fonts = (
        f'<w:rFonts w:ascii="{font_name}" w:hAnsi="{font_name}" '
        f'w:eastAsia="{font_name}" w:cs="{font_name}"/>'
        if font_name
        else ""
    )
    bold_xml = "" if not bold else "<w:b/>"
    italic_xml = "" if not italic else "<w:i/>"
    escaped = (
        text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )
    paragraphs_xml = "".join(
        f"<w:p><w:pPr><w:spacing w:before=\"0\" w:after=\"0\" w:line=\"240\" "
        f"w:lineRule=\"auto\"/><w:jc w:val=\"{alignment_value}\"/></w:pPr>"
        f"<w:r><w:rPr>{fonts}{bold_xml}{italic_xml}"
        f'<w:color w:val="{color_hex}"/><w:sz w:val="{font_size_half_points}"/>'
        f'<w:szCs w:val="{font_size_half_points}"/></w:rPr>'
        f"<w:t xml:space=\"preserve\">{line}</w:t></w:r></w:p>"
        for line in (escaped.split("\n") if escaped else [""])
    )
    rotation_attr = f' rot="{int(round(rotation * 60000))}"' if rotation else ""
    xml = (
        f'<w:r xmlns:w="{_WML_NS}" xmlns:wp="{_WP_NS}" xmlns:a="{_A_NS}" '
        f'xmlns:mc="{_MC_NS}" xmlns:wps="{_WPS_NS}" xmlns:v="{_V_NS}" '
        f'xmlns:w10="{_W10_NS}">'
        f'<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>'
        f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="{max(1, 251658240 + int(z_order) * 100)}" behindDoc="0" '
        f'locked="0" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{points_to_emu(x_points)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{points_to_emu(y_points)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{width_emu}" cy="{height_emu}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        f'<wp:docPr id="{max(1, int(object_id))}" name="{name or "TextBox"}"/>'
        f'<wp:cNvGraphicFramePr/>'
        f"<a:graphic><a:graphicData uri=\"{_WPS_NS}\">"
        f'<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr>'
        f'<a:xfrm{rotation_attr}><a:off x="0" y="0"/><a:ext cx="{width_emu}" cy="{height_emu}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f"<a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>"
        f"<wps:txbx><w:txbxContent>{paragraphs_xml}</w:txbxContent></wps:txbx>"
        f'<wps:bodyPr rot="0" wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" '
        f'anchor="t" anchorCtr="0"><a:noAutofit/></wps:bodyPr>'
        f"</wps:wsp></a:graphicData></a:graphic>"
        f"</wp:anchor></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict><v:shape id=\"{box_name}\" "
        f'style="position:absolute;left:{x_points}pt;top:{y_points}pt;'
        f'width:{width_points}pt;height:{height_points}pt" '
        f'fillcolor="none" stroked="f"><v:textbox inset="0,0,0,0">'
        f"<w:txbxContent>{paragraphs_xml}</w:txbxContent>"
        f"</v:textbox></v:shape></w:pict></mc:Fallback>"
        f"</mc:AlternateContent></w:r>"
    ).encode("utf-8")
    target_paragraph._p.append(parse_xml(xml))
    return target_paragraph


def position_table(
    table: Any,
    *,
    x_points: float,
    y_points: float,
    z_order: int = 0,
    overlap: bool = True,
    horz_anchor: str = "page",
    vert_anchor: str = "page",
) -> Any:
    """用 ``w:tblpPr`` 把表格浮动定位到页面坐标。"""
    table_properties = table._tbl.tblPr
    existing = table_properties.find(qn("w:tblpPr"))
    if existing is not None:
        table_properties.remove(existing)
    positioning = OxmlElement("w:tblpPr")
    positioning.set(qn("w:leftFromText"), "0")
    positioning.set(qn("w:rightFromText"), "0")
    positioning.set(qn("w:topFromText"), "0")
    positioning.set(qn("w:bottomFromText"), "0")
    positioning.set(qn("w:vertAnchor"), vert_anchor)
    positioning.set(qn("w:horzAnchor"), horz_anchor)
    positioning.set(qn("w:tblpX"), str(points_to_twips(x_points)))
    positioning.set(qn("w:tblpY"), str(points_to_twips(y_points)))
    style = table_properties.find(qn("w:tblStyle"))
    if style is not None:
        style.addnext(positioning)
    else:
        table_properties.insert(0, positioning)
    overlap_element = table_properties.find(qn("w:tblOverlap"))
    if overlap_element is None:
        overlap_element = OxmlElement("w:tblOverlap")
        positioning.addnext(overlap_element)
    overlap_element.set(qn("w:val"), "overlap" if overlap else "never")
    return table


def prepare_header(section: Any) -> Any:
    """确保 section 拥有独立的 header part，并返回可写入的容器。"""
    header = section.header
    try:
        header.is_linked_to_previous = False
    except Exception:
        pass
    return header


def prepare_footer(section: Any) -> Any:
    """确保 section 拥有独立的 footer part，并返回可写入的容器。"""
    footer = section.footer
    try:
        footer.is_linked_to_previous = False
    except Exception:
        pass
    return footer


def detach_header_footer(section: Any) -> None:
    """让本节拥有独立的空白页眉页脚，避免继承上一节的页码等内容。"""
    for part in (section.header, section.footer):
        try:
            part.is_linked_to_previous = False
        except Exception:
            continue
        try:
            for paragraph in list(part.paragraphs):
                clear_paragraph(paragraph)
                make_flow_paragraph_minimal(paragraph)
        except Exception:
            continue


def header_footer_paragraph(container: Any) -> Any:
    """返回 header/footer 中可复用的首个空段落。"""
    paragraphs = list(container.paragraphs)
    if paragraphs:
        paragraph = paragraphs[0]
        clear_paragraph(paragraph)
        return paragraph
    return container.add_paragraph()
