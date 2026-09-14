"""字体度量标定：用 Word 实测 framePr 原点到字形框顶部的偏移。

不同字体在 Word 的 EXACT 行距下，基线位置不同（例如内嵌 SimSun 与
系统 Arial 差 1pt 以上）。导出高保真 DOCX 前，先用一个小样本页把每种
字体的偏移系数测出来：

    k = (glyph_top - frame_top) / font_size

导出时令 ``frame_y = source_glyph_top - k * font_size`` 即可对齐。
结果按字体数据哈希缓存，重复转换不需要重新渲染。
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from .embedding import EmbeddedFontProgram, attach_embedded_fonts, is_cjk_font


DEFAULT_OFFSET_FACTOR = 0.074
"""未标定字体的经验偏移系数（Latin/系统字体），沿用历史实现。"""

CJK_SAMPLE = "字体标定黑体宋体排版测试样本内容"
LATIN_SAMPLE = "Hxg FontCalibration Sample 0123456789"
CALIBRATION_SIZES = (9.0, 10.5, 12.0, 16.0, 22.0)
_LINE_SPACING_POINTS = 46.0
_PAGE_HEIGHT_POINTS = 842.0
_PAGE_WIDTH_POINTS = 595.0
_TWIPS_PER_POINT = 20


@dataclass(frozen=True)
class FontMetricRequest:
    """需要标定的一个字体。"""

    word_name: str
    family: str
    sample: str = ""
    program: EmbeddedFontProgram | None = None

    @property
    def cache_key(self) -> str:
        digest = ""
        if self.program is not None:
            digest = hashlib.sha1(self.program.data).hexdigest()
        sample_class = "cjk" if _looks_cjk(self.family, self.sample) else "latin"
        return f"{self.word_name}|{sample_class}|{digest[:16]}"


@dataclass
class FontMetricTable:
    """字体 -> 偏移系数。"""

    offsets: dict[str, float] = field(default_factory=dict)
    sources: dict[str, str] = field(default_factory=dict)

    def factor(self, word_name: str) -> float | None:
        return self.offsets.get(word_name)

    def to_dict(self) -> dict[str, Any]:
        return {
            "offsets": {key: round(value, 5) for key, value in self.offsets.items()},
            "sources": dict(self.sources),
        }


def _looks_cjk(family: str, sample: str = "") -> bool:
    if is_cjk_font(family):
        return True
    return bool(re.search(r"[\u3400-\u4dbf\u4e00-\u9fff]", sample or ""))


def default_cache_path() -> Path:
    import os

    configured = os.getenv("PDF_FONT_METRICS_CACHE", "").strip()
    if configured:
        return Path(configured)
    return (
        Path(__file__).resolve().parents[2] / "model_cache" / "font_metrics.json"
    )


def load_cache(path: Path | None = None) -> dict[str, float]:
    target = Path(path) if path else default_cache_path()
    if not target.is_file():
        return {}
    try:
        payload = json.loads(target.read_text(encoding="utf-8"))
    except Exception:
        return {}
    cached = payload.get("offsets") if isinstance(payload, dict) else None
    if isinstance(cached, dict):
        return {
            str(key): float(value)
            for key, value in cached.items()
            if isinstance(value, (int, float))
        }
    return {}


def save_cache(offsets: dict[str, float], path: Path | None = None) -> None:
    target = Path(path) if path else default_cache_path()
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        existing = load_cache(target)
        existing.update({key: value for key, value in offsets.items()})
        target.write_text(
            json.dumps({"offsets": existing}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    except Exception:
        return


# ------------------------------------------------------------------ 生成样本
def _add_calibration_paragraph(
    document: Any,
    *,
    word_name: str,
    size: float,
    y: float,
    sample: str,
) -> None:
    paragraph = document.add_paragraph()
    properties = paragraph._p.get_or_add_pPr()
    frame = OxmlElement("w:framePr")
    frame.set(qn("w:wrap"), "none")
    frame.set(qn("w:hAnchor"), "page")
    frame.set(qn("w:vAnchor"), "page")
    frame.set(qn("w:x"), str(int(25 * _TWIPS_PER_POINT)))
    frame.set(qn("w:y"), str(int(round(y * _TWIPS_PER_POINT))))
    frame.set(qn("w:w"), str(int(540 * _TWIPS_PER_POINT)))
    frame.set(qn("w:hRule"), "auto")
    properties.insert(0, frame)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(size)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    run = paragraph.add_run(sample)
    run.font.size = Pt(size)
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), word_name)


def build_calibration_document(
    requests: Iterable[FontMetricRequest],
) -> tuple[Any, list[dict[str, Any]]]:
    """生成标定 DOCX 与每个样本的期望位置。"""
    document = Document()
    _reset_section(document.sections[0])
    entries: list[dict[str, Any]] = []
    y = 30.0
    page_index = 0
    for request in requests:
        sample = request.sample or (
            CJK_SAMPLE if _looks_cjk(request.family) else LATIN_SAMPLE
        )
        for size in CALIBRATION_SIZES:
            if y > _PAGE_HEIGHT_POINTS - 60.0:
                section = document.add_section()
                _reset_section(section)
                y = 30.0
                page_index += 1
            _add_calibration_paragraph(
                document,
                word_name=request.word_name,
                size=size,
                y=y,
                sample=sample,
            )
            entries.append(
                {
                    "word_name": request.word_name,
                    "family": request.family,
                    "size": size,
                    "y": y,
                    "sample": sample,
                    "page": page_index,
                }
            )
            y += _LINE_SPACING_POINTS
    programs = [
        request.program
        for request in requests
        if request.program is not None
    ]
    if programs:
        attach_embedded_fonts(document, programs)
    return document, entries


def _reset_section(section: Any) -> None:
    section.page_width = Pt(_PAGE_WIDTH_POINTS)
    section.page_height = Pt(_PAGE_HEIGHT_POINTS)
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, margin, Pt(0))


# ------------------------------------------------------------------ 测量
def measure_offsets(
    pdf_path: Path,
    entries: list[dict[str, Any]],
) -> dict[str, list[float]]:
    """从渲染出的 PDF 中测量每个样本的字形框顶部，返回 字体 -> k 列表。"""
    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_raw

    samples: dict[str, list[float]] = {}
    used: set[int] = set()
    document = pdfium.PdfDocument(str(pdf_path))
    try:
        for page_index in range(len(document)):
            page = document[page_index]
            text_page = page.get_textpage()
            try:
                characters: list[tuple[float, float, float, float]] = []
                try:
                    count = int(text_page.count_chars())
                except Exception:
                    count = len(text_page.get_text_range())
                for index in range(max(count, 0)):
                    try:
                        codepoint = int(
                            pdfium_raw.FPDFText_GetUnicode(
                                text_page.raw,
                                index,
                            )
                        )
                        character = (
                            chr(codepoint)
                            if codepoint and codepoint <= 0x10FFFF
                            else ""
                        )
                    except Exception:
                        try:
                            character = text_page.get_text_range(index, 1)[:1]
                        except Exception:
                            try:
                                character = text_page.get_text_range()[
                                    index : index + 1
                                ]
                            except Exception:
                                character = ""
                    if character in {"", "\r", "\n"}:
                        continue
                    try:
                        x0, y0, x1, y1 = text_page.get_charbox(index)
                    except Exception:
                        continue
                    if x1 - x0 <= 0 and y1 - y0 <= 0:
                        continue
                    characters.append(
                        (float(x0), float(y0), float(x1), float(y1))
                    )
            finally:
                text_page.close()
                page.close()
            lines: list[list[tuple[float, float, float, float]]] = []
            for item in sorted(characters, key=lambda value: (-value[3], value[0])):
                placed = False
                for line in lines:
                    if abs(item[3] - line[0][3]) <= 5.0:
                        line.append(item)
                        placed = True
                        break
                if not placed:
                    lines.append([item])
            for line in lines:
                top = _PAGE_HEIGHT_POINTS - max(item[3] for item in line)
                _match_entry(samples, entries, page_index, top, used)
    finally:
        document.close()
    return samples


def _match_entry(
    samples: dict[str, list[float]],
    entries: list[dict[str, Any]],
    page_index: int,
    top: float,
    used: set[int],
) -> None:
    """把一条检测到的文本行匹配到最近的标定条目（每条目只用一次）。"""
    best = None
    for position, entry in enumerate(entries):
        if position in used:
            continue
        if entry.get("page") not in (None, page_index):
            continue
        distance = abs(top - entry["y"])
        window = max(0.35 * float(entry["size"]), 5.0)
        if distance > window:
            continue
        if best is None or distance < best[0]:
            best = (distance, position, entry)
    if best is None:
        return
    used.add(best[1])
    entry = best[2]
    factor = (top - entry["y"]) / entry["size"]
    samples.setdefault(entry["word_name"], []).append(factor)


# ------------------------------------------------------------------ 入口
def calibrate_font_metrics(
    requests: Iterable[FontMetricRequest],
    *,
    work_dir: Path,
    cache_path: Path | None = None,
    timeout_seconds: float = 180.0,
    renderer: Any = None,
    use_cache: bool = True,
) -> FontMetricTable:
    """标定字体偏移；命中缓存时不渲染。"""
    request_list = [request for request in requests if request.word_name]
    table = FontMetricTable()
    cached = load_cache(cache_path) if use_cache else {}
    pending: list[FontMetricRequest] = []
    for request in request_list:
        key = request.cache_key
        if key in cached:
            table.offsets[request.word_name] = cached[key]
            table.sources[request.word_name] = "cache"
        else:
            pending.append(request)
    if not pending:
        return table

    work_dir = Path(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    docx_path = work_dir / "font-metrics.docx"
    pdf_path = work_dir / "font-metrics.pdf"
    document, entries = build_calibration_document(pending)
    document.save(str(docx_path))
    if renderer is None:
        from ..validate.render import render_docx_to_pdf

        renderer = render_docx_to_pdf
    renderer(docx_path, pdf_path, timeout_seconds=timeout_seconds)

    measured = measure_offsets(pdf_path, entries)
    updated: dict[str, float] = {}
    for request in pending:
        values = measured.get(request.word_name) or []
        if not values:
            continue
        factor = round(sum(values) / len(values), 5)
        table.offsets[request.word_name] = factor
        table.sources[request.word_name] = f"measured:{len(values)}"
        updated[request.cache_key] = factor
    if updated:
        save_cache(updated, cache_path)
    return table
