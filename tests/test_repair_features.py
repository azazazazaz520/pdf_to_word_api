"""修复方案中新增的区域、坐标、归属和独立质量检查。"""

from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest

from docx import Document
from reportlab.pdfgen.canvas import Canvas

from src.layout.models import PdfLayoutRegion, PdfPageLayout, PdfTextLine
from src.layout.regions import (
    PageCoordinateTransform,
    assign_model_regions_to_lines,
    merge_local_ocr_lines,
    normalize_model_regions,
)
from src.layout.text import _matrix_scale, _matrix_scales
from src.layout.vectors import _extract_vector_lines
from src.layout.reading_order import _mark_running_headers
from src.ooxml_positioning import add_absolute_text_box, new_canvas_paragraph
from src.pdf_worker import _quality_source_block_text
from src.worker.quality import (
    _apply_fidelity_acceptance,
    _evaluate_quality_gate,
    evaluate_final_content_quality,
    formula_text_exception_is_local,
    read_docx_text,
)
import pypdfium2 as pdfium


class RepairFeatureTest(unittest.TestCase):
    def test_quality_source_text_uses_export_order_for_tables_and_bullets(self) -> None:
        table_block = SimpleNamespace(
            kind="table",
            table=SimpleNamespace(
                rows=(("12e", ""), ("13a", "")),
                cells=(SimpleNamespace(text="13a"), SimpleNamespace(text="12e")),
            ),
            lines=(),
            text="13a12e",
        )
        bullet_block = SimpleNamespace(
            kind="bullet",
            table=None,
            lines=(
                SimpleNamespace(text="• Single or Married filing separately"),
                SimpleNamespace(text="• Married filing jointly"),
            ),
            text="Single or Married filing separately Married filing jointly",
        )

        self.assertEqual(_quality_source_block_text(table_block), "12e\n13a")
        self.assertEqual(
            _quality_source_block_text(bullet_block),
            "• Single or Married filing separately\n• Married filing jointly",
        )

    def test_font_matrix_uses_height_scale_and_keeps_non_uniform_components(self) -> None:
        self.assertAlmostEqual(_matrix_scale((3, 0, 0, 4, 20, 20)), 4.0)
        self.assertEqual(_matrix_scales((3, 0, 0, 4, 20, 20)), (3.0, 4.0))

    def test_model_bbox_is_converted_with_rotation(self) -> None:
        transform = PageCoordinateTransform(
            pixel_width=200,
            pixel_height=100,
            page_width=400,
            page_height=200,
        )
        self.assertEqual(transform.bbox_to_page((10, 20, 110, 70)), (20.0, 40.0, 220.0, 140.0))
        rotated = PageCoordinateTransform(
            pixel_width=200,
            pixel_height=100,
            page_width=400,
            page_height=200,
            rotation=90,
        )
        self.assertEqual(rotated.bbox_to_page((0, 0, 200, 100)), (0.0, 0.0, 400.0, 200.0))

    def test_transformed_path_lines_use_page_coordinates(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            pdf_path = Path(temporary_directory) / "transformed-lines.pdf"
            canvas = Canvas(str(pdf_path), pagesize=(600, 600))
            canvas.saveState()
            canvas.translate(135, 100)
            canvas.rect(0, 0, 378, 40, stroke=1, fill=0)
            canvas.restoreState()
            canvas.save()

            document = pdfium.PdfDocument(str(pdf_path))
            try:
                horizontal, vertical = _extract_vector_lines(document[0], 600.0)
            finally:
                document.close()
            self.assertTrue(any(abs(line.x0 - 135.0) < 0.1 and abs(line.x1 - 513.0) < 0.1 for line in horizontal))
            self.assertTrue(any(abs(line.x - 135.0) < 0.1 for line in vertical))

    def test_model_regions_keep_page_source_and_order_fields(self) -> None:
        line = PdfTextLine(
            text="body",
            x0=20,
            top=20,
            x1=80,
            bottom=32,
            font_size=10,
            source_char_ids=("1:4",),
        )
        regions = normalize_model_regions(
            [{"bbox": [10, 10, 100, 50], "label": "text", "score": 0.8, "order_hint": 2}],
            page_number=1,
            transform=PageCoordinateTransform(100, 100, 200, 200),
            lines=[line],
        )
        self.assertEqual(regions[0].source, "layout-model")
        self.assertEqual(regions[0].order_hint, 2)
        self.assertEqual(regions[0].source_char_ids, ("1:4",))

    def test_model_regions_constrain_native_lines_to_columns(self) -> None:
        lines = (
            PdfTextLine("left", 30, 20, 80, 30, 10),
            PdfTextLine("right", 120, 20, 175, 30, 10),
        )
        regions = normalize_model_regions(
            [
                {"coordinate": [0, 0, 90, 100], "label": "text", "score": 0.9},
                {"coordinate": [100, 0, 190, 100], "label": "text", "score": 0.9},
            ],
            page_number=1,
            transform=PageCoordinateTransform(200, 100, 200, 100),
            lines=lines,
        )
        assigned, columns = assign_model_regions_to_lines(
            lines,
            model_regions=regions,
            page_width=200,
        )
        self.assertEqual(len(columns), 1)
        self.assertEqual([line.region_id for line in assigned], ["column-0", "column-1"])

    def test_local_ocr_does_not_duplicate_usable_native_text(self) -> None:
        region = PdfLayoutRegion("r1", 1, "body", (0, 0, 100, 100))
        native = PdfTextLine("Hello", 10, 10, 50, 20, 10)
        ocr = PdfTextLine("Hello", 10, 10, 50, 20, 10)
        merged, report = merge_local_ocr_lines([native], [ocr], region=region)
        self.assertEqual([line.text for line in merged], ["Hello"])
        self.assertEqual(report["source"], "native")

    def test_final_quality_gate_reports_missing_and_duplicate_source_chars(self) -> None:
        report = evaluate_final_content_quality(
            source_char_ids=("1:1", "1:2"),
            placements=[
                {"source_char_ids": ["1:1", "1:1"]},
            ],
            source_text="AB",
            output_text="A",
        )
        self.assertEqual(report["status"], "failed")
        self.assertEqual(
            {check["name"] for check in report["checks"] if check["status"] == "failed"},
            {
                "source_character_coverage",
                "source_character_unique_placement",
                "text_content_presence",
            },
        )

    def test_lost_word_space_fails_content_gate(self) -> None:
        report = evaluate_final_content_quality(
            source_char_ids=("1:1", "1:2"),
            placements=[{"source_char_ids": ["1:1", "1:2"], "status": "native"}],
            source_text="alpha beta",
            output_text="alphabeta",
        )
        self.assertEqual(report["status"], "failed")

    def test_duplicate_visible_text_fails_content_gate(self) -> None:
        report = evaluate_final_content_quality(
            source_char_ids=("1:1", "1:2"),
            placements=[{"source_char_ids": ["1:1", "1:2"], "status": "native"}],
            source_text="AB",
            output_text="ABAB",
        )
        self.assertEqual(report["status"], "failed")

    def test_reordered_visible_text_is_unverified(self) -> None:
        report = evaluate_final_content_quality(
            source_char_ids=("1:1", "1:2", "1:3"),
            placements=[{"source_char_ids": ["1:1", "1:2", "1:3"]}],
            source_text="alpha beta gamma",
            output_text="gamma alpha beta",
        )
        self.assertEqual(report["status"], "unverified")

    def test_whitespace_only_layout_difference_passes_content_gate(self) -> None:
        report = evaluate_final_content_quality(
            source_char_ids=("1:1", "1:2", "1:3"),
            placements=[{"source_char_ids": ["1:1", "1:2", "1:3"]}],
            source_text="alphabeta gamma",
            output_text="alpha beta gamma",
        )
        self.assertEqual(report["status"], "passed")
        self.assertIn(
            "source_sequence_matched_with_whitespace_variations",
            report["checks"][-1]["detail"],
        )

    def test_formula_exception_does_not_hide_ordinary_text_loss(self) -> None:
        blocks = (
            SimpleNamespace(kind="paragraph", block_id="p1", text="alpha"),
            SimpleNamespace(kind="formula", block_id="f1", text="x = y"),
            SimpleNamespace(kind="paragraph", block_id="p2", text="beta gamma"),
        )
        placement = {
            "block_id": "f1",
            "status": "native",
            "reason": "flow_formula_omml",
        }
        self.assertTrue(
            formula_text_exception_is_local(
                source_blocks=blocks,
                placements=[placement],
                output_text="alpha x = y beta gamma",
            )
        )
        self.assertFalse(
            formula_text_exception_is_local(
                source_blocks=blocks,
                placements=[placement],
                output_text="alpha x = y gamma",
            )
        )

    def test_docx_readback_includes_referenced_header_and_footer(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "readback.docx"
            document = Document()
            document.add_paragraph("正文")
            document.sections[0].header.paragraphs[0].text = "页眉证据"
            document.sections[0].footer.paragraphs[0].text = "页脚证据"
            document.save(output_path)

            text = read_docx_text(output_path)
            body_text = read_docx_text(
                output_path,
                include_headers_footers=False,
            )

        self.assertIn("正文", text)
        self.assertIn("页眉证据", text)
        self.assertIn("页脚证据", text)
        self.assertEqual(body_text, "正文")

    def test_docx_readback_does_not_double_count_positioned_text_box(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "positioned-readback.docx"
            document = Document()
            canvas = new_canvas_paragraph(document)
            add_absolute_text_box(
                document,
                x_points=10,
                y_points=10,
                width_points=100,
                height_points=20,
                text="定位文本",
                paragraph=canvas,
            )
            document.save(output_path)

            text = read_docx_text(output_path)

        self.assertEqual(text, "定位文本")

    def test_same_page_dots_do_not_become_running_footer(self) -> None:
        dots = (
            PdfTextLine(".", 10, 720, 11, 722, 8),
            PdfTextLine(".", 20, 720, 21, 722, 8),
        )
        pages = (
            PdfPageLayout(600, 800, dots, (), ()),
            PdfPageLayout(600, 800, (), (), ()),
        )
        result = _mark_running_headers(pages)
        self.assertFalse(any(line.is_header_footer for line in result[0].lines))

    def test_form_field_labels_at_different_edges_do_not_become_running_text(self) -> None:
        pages = (
            PdfPageLayout(
                600,
                800,
                (PdfTextLine("7a", 40, 700, 55, 710, 8),),
                (),
                (),
            ),
            PdfPageLayout(
                600,
                800,
                (PdfTextLine("12a", 40, 40, 65, 50, 8),),
                (),
                (),
            ),
        )
        result = _mark_running_headers(pages)
        self.assertFalse(any(line.is_header_footer for page in result for line in page.lines))

    def test_missing_visual_evidence_is_unverified(self) -> None:
        render = {
            "status": "succeeded",
            "source_page_count": 1,
            "rendered_page_count": 1,
            "page_delta": 0,
            "unexpected_blank_pages": [],
            "ssim": {"status": "not_run"},
            "text_layout": {"status": "not_run"},
        }
        quality = {"export_mode": "fidelity", "render_validation": render}
        _apply_fidelity_acceptance(quality, render, ssim_threshold=0.98)
        report = _evaluate_quality_gate(
            quality,
            enabled=True,
            page_delta_warn_ratio=0.05,
            page_delta_warn_absolute=3,
        )
        self.assertEqual(report["status"], "unverified")


if __name__ == "__main__":
    unittest.main()
