from __future__ import annotations

import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from docx import Document
from PIL import Image

from src.export.structured import (
    _complex_table_regions,
    _dense_form_background_blocks,
    _positioned_table_cell_metrics,
    _positioned_table_cell_lines,
    export_structured_docx,
)
from src.ir.model import IRBlock, IRDocument, IRPage, IRTextLine
from src.layout.models import PdfTable, PdfTableCell


def _image_bytes() -> bytes:
    output = BytesIO()
    Image.new("RGB", (120, 60), color=(40, 100, 180)).save(output, format="PNG")
    return output.getvalue()


def _pdf_table() -> PdfTable:
    return PdfTable(
        bbox=(40.0, 70.0, 320.0, 130.0),
        column_boundaries=(40.0, 180.0, 320.0),
        row_boundaries=(70.0, 100.0, 130.0),
        rows=(("表头 A", "表头 B"), ("单元格 1", "单元格 2")),
        cells=(
            PdfTableCell(
                row_index=0,
                column_index=0,
                row_span=1,
                column_span=1,
                bbox=(40.0, 70.0, 180.0, 100.0),
                text="表头 A",
                font_size=9.0,
                bold=True,
            ),
            PdfTableCell(
                row_index=0,
                column_index=1,
                row_span=1,
                column_span=1,
                bbox=(180.0, 70.0, 320.0, 100.0),
                text="表头 B",
                font_size=9.0,
                bold=True,
            ),
            PdfTableCell(
                row_index=1,
                column_index=0,
                row_span=1,
                column_span=1,
                bbox=(40.0, 100.0, 180.0, 130.0),
                text="单元格 1",
                font_size=9.0,
            ),
            PdfTableCell(
                row_index=1,
                column_index=1,
                row_span=1,
                column_span=1,
                bbox=(180.0, 100.0, 320.0, 130.0),
                text="单元格 2",
                font_size=9.0,
            ),
        ),
    )


def _complex_pdf_table() -> PdfTable:
    return PdfTable(
        bbox=(40.0, 70.0, 320.0, 120.0),
        column_boundaries=(40.0, 180.0, 320.0),
        row_boundaries=(70.0, 82.0, 120.0),
        rows=(("A\nB\nC", "1\n2\n3"), ("D", "4")),
        cells=(
            PdfTableCell(
                row_index=0,
                column_index=0,
                row_span=1,
                column_span=1,
                bbox=(40.0, 70.0, 180.0, 82.0),
                text="A\nB\nC",
                font_size=10.0,
            ),
            PdfTableCell(
                row_index=0,
                column_index=1,
                row_span=1,
                column_span=1,
                bbox=(180.0, 70.0, 320.0, 82.0),
                text="1\n2\n3",
                font_size=10.0,
            ),
        ),
    )


class StructuredExportTest(unittest.TestCase):
    def test_dense_form_background_selection_keeps_only_large_color_fills(self) -> None:
        page = IRPage(
            page_number=1,
            width=600.0,
            height=800.0,
            route="text",
            blocks=[
                IRBlock(
                    kind="vector",
                    page=1,
                    bbox=(40.0, 20.0, 560.0, 50.0),
                    vector=SimpleNamespace(
                        fill_color=(220, 255, 251),
                        stroke_color=None,
                    ),
                ),
                IRBlock(
                    kind="vector",
                    page=1,
                    bbox=(480.0, 100.0, 500.0, 112.0),
                    vector=SimpleNamespace(
                        fill_color=(192, 192, 192),
                        stroke_color=None,
                    ),
                ),
                IRBlock(
                    kind="vector",
                    page=1,
                    bbox=(40.0, 60.0, 560.0, 90.0),
                    vector=SimpleNamespace(
                        fill_color=(255, 255, 255),
                        stroke_color=None,
                    ),
                ),
            ],
        )

        backgrounds = _dense_form_background_blocks(page)

        self.assertEqual(len(backgrounds), 1)
        self.assertEqual(backgrounds[0].bbox, (40.0, 20.0, 560.0, 50.0))

    def test_positioned_table_cell_keeps_space_glyph_on_text_line(self) -> None:
        glyphs = (
            SimpleNamespace(text="I", bbox=(0.0, 0.0, 1.0, 5.0), font_size=7.5),
            SimpleNamespace(text="f", bbox=(1.2, 0.0, 2.2, 5.0), font_size=7.5),
            SimpleNamespace(text=" ", bbox=(2.3, 4.8, 3.3, 4.81), font_size=7.5),
            SimpleNamespace(text="m", bbox=(3.5, 1.0, 5.5, 5.0), font_size=7.5),
            SimpleNamespace(text="o", bbox=(5.7, 1.0, 7.0, 5.0), font_size=7.5),
            SimpleNamespace(text="r", bbox=(7.2, 1.0, 8.2, 5.0), font_size=7.5),
            SimpleNamespace(text="e", bbox=(8.4, 1.0, 9.7, 5.0), font_size=7.5),
        )
        lines = _positioned_table_cell_lines(SimpleNamespace(glyphs=glyphs))

        self.assertEqual([text for text, _, _ in lines], ["If more"])

    def test_positioned_table_cell_box_leaves_word_text_room(self) -> None:
        metrics = _positioned_table_cell_metrics(
            SimpleNamespace(bbox=(10.0, 10.0, 60.0, 40.0), font_size=7.0),
            "label",
            (12.0, 15.0, 48.0, 20.0),
            7.0,
        )

        self.assertGreaterEqual(metrics[2], 37.0)
        self.assertGreaterEqual(metrics[3], 8.4)
        self.assertLessEqual(metrics[0] + metrics[2], 60.0)

    def test_overfull_table_cells_are_marked_for_region_fallback(self) -> None:
        page = IRPage(
            page_number=1,
            width=400.0,
            height=300.0,
            route="text",
            blocks=[
                IRBlock(
                    kind="table",
                    page=1,
                    bbox=(40.0, 70.0, 320.0, 120.0),
                    table=_complex_pdf_table(),
                )
            ],
        )

        self.assertEqual(
            _complex_table_regions(page),
            ((38.0, 68.0, 322.0, 122.0),),
        )

    def test_line_spacing_ratio_is_converted_to_points(self) -> None:
        ir = IRDocument(
            pages=[
                IRPage(
                    page_number=1,
                    width=400.0,
                    height=300.0,
                    route="text",
                    blocks=[
                        IRBlock(
                            kind="paragraph",
                            page=1,
                            text="第一行 第二行",
                            font_size=10.0,
                            line_spacing=1.1,
                            lines=(
                                IRTextLine(
                                    text="第一行",
                                    bbox=(40.0, 40.0, 80.0, 50.0),
                                    font_size=10.0,
                                ),
                                IRTextLine(
                                    text="第二行",
                                    bbox=(40.0, 51.0, 80.0, 61.0),
                                    font_size=10.0,
                                ),
                            ),
                        )
                    ],
                )
            ]
        )

        with TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "line-spacing.docx"
            export_structured_docx(ir, output_path)
            with zipfile.ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml").decode("utf-8")

        self.assertIn('w:line="220"', document_xml)
        self.assertNotIn('w:line="22"', document_xml)

    def test_rotated_text_remains_editable_without_image_fallback(self) -> None:
        ir = IRDocument(
            pages=[
                IRPage(
                    page_number=1,
                    width=400.0,
                    height=300.0,
                    route="text",
                    blocks=[
                        IRBlock(
                            kind="paragraph",
                            page=1,
                            text="侧边元数据",
                            bbox=(8.0, 80.0, 20.0, 180.0),
                            rotation=-90.0,
                            fallback_image=_image_bytes(),
                        )
                    ],
                )
            ]
        )

        with TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "rotated-text.docx"
            report = export_structured_docx(ir, output_path)
            with zipfile.ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml").decode("utf-8")

        self.assertEqual(report["rotated_text_image_fallback_count"], 0)
        self.assertNotIn("<wp:anchor", document_xml)
        self.assertIn("侧边元数据", document_xml)

    def test_grid_like_text_remains_editable_without_image_fallback(self) -> None:
        lines = tuple(
            IRTextLine(
                text=text,
                bbox=bbox,
                font_size=10.0,
            )
            for text, bbox in (
                ("甲", (40.0, 40.0, 60.0, 50.0)),
                ("乙", (180.0, 40.0, 200.0, 50.0)),
                ("丙", (40.0, 55.0, 60.0, 65.0)),
                ("丁", (180.0, 55.0, 200.0, 65.0)),
            )
        )
        ir = IRDocument(
            pages=[
                IRPage(
                    page_number=1,
                    width=400.0,
                    height=300.0,
                    route="text",
                    blocks=[
                        IRBlock(
                            kind="paragraph",
                            page=1,
                            text="甲乙丙丁",
                            bbox=(40.0, 40.0, 200.0, 65.0),
                            lines=lines,
                            fallback_image=_image_bytes(),
                        )
                    ],
                )
            ]
        )

        with TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "grid-text.docx"
            report = export_structured_docx(ir, output_path)
            document = Document(output_path)
            with zipfile.ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml").decode("utf-8")

        self.assertEqual(report["text_image_fallback_count"], 0)
        self.assertEqual(len(document.paragraphs), 1)
        self.assertNotIn("ListBullet", document_xml)
        self.assertIn("甲乙丙丁", document_xml)

    def test_writes_flow_text_table_omml_and_inline_image(self) -> None:
        ir = IRDocument(
            title="结构化示例",
            metadata={"outline": [{"page": 1, "title": "Chapter 1"}]},
            pages=[
                IRPage(
                    page_number=1,
                    width=400.0,
                    height=300.0,
                    route="text",
                    blocks=[
                        IRBlock(
                            kind="paragraph",
                            page=1,
                            text="普通正文应当保持为可重排的 Word 段落。",
                        ),
                        IRBlock(
                            kind="table",
                            page=1,
                            bbox=(40.0, 70.0, 320.0, 130.0),
                            table=_pdf_table(),
                        ),
                        IRBlock(
                            kind="formula",
                            page=1,
                            text=r"x = \frac{a}{b} + 1",
                            bbox=(40.0, 140.0, 320.0, 170.0),
                        ),
                        IRBlock(
                            kind="image",
                            page=1,
                            bbox=(40.0, 180.0, 160.0, 240.0),
                            image_bytes=_image_bytes(),
                            image_width=120.0,
                            image_height=60.0,
                            image_alt="示例图片",
                        ),
                    ],
                )
            ],
        )

        with TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "result.docx"
            report = export_structured_docx(
                ir,
                output_path,
                include_toc=True,
                include_bookmarks=True,
            )

            self.assertEqual(report["text_paragraph_count"], 1)
            self.assertEqual(report["table_count"], 1)
            self.assertEqual(report["formula_count"], 1)
            self.assertEqual(report["formula_omml_count"], 1)
            self.assertEqual(report["image_count"], 1)
            document = Document(output_path)
            self.assertEqual(len(document.tables), 1)
            self.assertIn(
                "普通正文应当保持为可重排的 Word 段落。",
                "\n".join(paragraph.text for paragraph in document.paragraphs),
            )
            self.assertEqual(document.tables[0].cell(1, 1).text, "单元格 2")
            with zipfile.ZipFile(output_path) as archive:
                document_xml = archive.read("word/document.xml").decode("utf-8")
            self.assertNotIn("<w:framePr", document_xml)
            self.assertIn("<w:tbl", document_xml)
            self.assertIn("<m:oMath", document_xml)
            self.assertIn("<wp:inline", document_xml)
            self.assertIn("w:bookmarkStart", document_xml)
            self.assertIn("Chapter_1", document_xml)
            self.assertIn("TOC", document_xml)

    def test_writes_ocr_html_table_as_word_table(self) -> None:
        ir = IRDocument(
            pages=[
                IRPage(
                    page_number=1,
                    width=400.0,
                    height=300.0,
                    route="ocr",
                    blocks=[
                        IRBlock(
                            kind="html_table",
                            page=1,
                            text=(
                                "<table><tr><th colspan='2'>总览</th></tr>"
                                "<tr><td rowspan='2'>A</td><td>B</td></tr>"
                                "<tr><td>C</td></tr></table>"
                            ),
                        )
                    ],
                )
            ]
        )

        with TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "ocr-table.docx"
            report = export_structured_docx(ir, output_path)

            self.assertEqual(report["table_count"], 1)
            self.assertEqual(report["html_table_count"], 1)
            document = Document(output_path)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(document.tables[0].cell(0, 0).text, "总览")
            self.assertEqual(document.tables[0].cell(1, 0).text, "A")
            self.assertEqual(document.tables[0].cell(2, 1).text, "C")

    def test_bad_table_is_degraded_locally(self) -> None:
        ir = IRDocument(
            pages=[
                IRPage(
                    page_number=1,
                    width=400.0,
                    height=300.0,
                    route="text",
                    blocks=[IRBlock(kind="table", page=1, table=object())],
                )
            ]
        )

        with TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "bad-table.docx"
            report = export_structured_docx(ir, output_path)

            self.assertTrue(output_path.is_file())
            self.assertEqual(report["table_native_error_count"], 1)
            self.assertEqual(report["skipped_block_count"], 1)
            self.assertEqual(report["placements"][0]["status"], "skipped")

    def test_dense_vectors_are_compacted_without_per_object_rendering(self) -> None:
        ir = IRDocument(
            pages=[
                IRPage(
                    page_number=1,
                    width=400.0,
                    height=300.0,
                    route="text",
                    blocks=[
                        IRBlock(
                            kind="vector",
                            page=1,
                            bbox=(20.0 + index, 40.0, 21.0 + index, 41.0),
                        )
                        for index in range(65)
                    ],
                )
            ]
        )

        with TemporaryDirectory() as temporary_directory:
            output_path = Path(temporary_directory) / "dense-vectors.docx"
            report = export_structured_docx(ir, output_path)

            self.assertEqual(report["vector_compacted_page_count"], 1)
            self.assertEqual(report["vector_skipped_count"], 65)
            self.assertEqual(report["vector_image_fallback_count"], 0)
            self.assertTrue(output_path.is_file())


if __name__ == "__main__":
    unittest.main()
