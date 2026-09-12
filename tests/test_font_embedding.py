from __future__ import annotations

import unittest
import uuid
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from src.fonts.embedding import EmbeddedFontProgram, _looks_like_sfnt, attach_embedded_fonts, build_font_plan, extract_embedded_font_programs, guid_bytes, is_cjk_font, obfuscate_font_data
from src.fonts.metrics import CALIBRATION_SIZES, FontMetricRequest, build_calibration_document, measure_offsets


FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "synthetic_text_table.pdf"


class FontEmbeddingTest(unittest.TestCase):
    def test_guid_key_is_reversed_hex(self) -> None:
        # 用 Word 自己生成的嵌入字体反推出的顺序：GUID 十六进制字节整体反转
        guid = "{3F9D2B2E-BDDB-4393-9B69-21D430E5E38A}"
        self.assertEqual(
            guid_bytes(guid).hex(),
            "8ae3e530d421699b9343dbbd2e2b9d3f",
        )

    def test_obfuscation_round_trip(self) -> None:
        data = bytes(range(64))
        guid = "{" + str(uuid.uuid4()).upper() + "}"
        masked = obfuscate_font_data(data, guid)
        self.assertNotEqual(masked[:32], data[:32])
        self.assertEqual(masked[32:], data[32:])
        self.assertEqual(obfuscate_font_data(masked, guid), data)

    def test_extract_and_plan_for_real_pdf(self) -> None:
        if not FIXTURE.is_file():
            self.skipTest("缺少合成样本")
        programs = extract_embedded_font_programs(FIXTURE)
        self.assertTrue(programs)
        sample = next(
            (item for item in programs.values() if "YaHei" in item.family),
            None,
        )
        self.assertIsNotNone(sample)
        self.assertTrue(_looks_like_sfnt(sample.data))
        self.assertTrue(is_cjk_font(sample.family, sample.raw_name))

        plan = build_font_plan(
            FIXTURE,
            [
                {
                    "pdf_font_name": sample.raw_name,
                    "font_name": sample.family,
                    "bold": False,
                    "italic": False,
                },
                {
                    "pdf_font_name": "Helvetica",
                    "font_name": "Arial",
                    "bold": False,
                    "italic": False,
                },
            ],
        )
        report = plan.report()
        self.assertEqual(report["embedded_font_count"], 1)
        self.assertIn("Helvetica", report["fallback_fonts"])
        entry = plan.entry(sample.raw_name)
        self.assertIsNotNone(entry)
        self.assertTrue(entry.embedded)
        self.assertEqual(plan.word_name_for(sample.raw_name, "fallback"), entry.word_name)
        # 内嵌字体自带粗体字形，不再叠加 w:b
        self.assertEqual(plan.style_for(sample.raw_name, True, True), (False, False))
        self.assertEqual(plan.style_for("Helvetica", True, False), (True, False))

    def test_attach_embedded_fonts_writes_package_parts(self) -> None:
        if not FIXTURE.is_file():
            self.skipTest("缺少合成样本")
        programs = extract_embedded_font_programs(FIXTURE)
        program = next(iter(programs.values()))
        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "embedded.docx"
            document = Document()
            run = document.add_paragraph().add_run("嵌入字体测试")
            run.font.size = Pt(12)
            run.font.name = program.word_name
            properties = run._element.get_or_add_rPr()
            fonts = properties.get_or_add_rFonts()
            for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fonts.set(qn(attribute), program.word_name)
            ids = attach_embedded_fonts(document, [program])
            self.assertIn(program.word_name, ids)
            document.save(output)

            with zipfile.ZipFile(output) as archive:
                names = archive.namelist()
                font_parts = [name for name in names if name.endswith(".odttf")]
                self.assertEqual(len(font_parts), 1)
                table = archive.read("word/fontTable.xml").decode("utf-8")
                settings = archive.read("word/settings.xml").decode("utf-8")
                rels = archive.read("word/_rels/fontTable.xml.rels").decode("utf-8")
                content_types = archive.read("[Content_Types].xml").decode("utf-8")
            self.assertIn("w:embedRegular", table)
            self.assertIn(program.guid, table)
            self.assertIn("embedTrueTypeFonts", settings)
            self.assertIn("relationships/font", rels)
            self.assertIn("obfuscatedFont", content_types)


class FontMetricsTest(unittest.TestCase):
    def test_calibration_document_layout(self) -> None:
        program = None
        if FIXTURE.is_file():
            programs = extract_embedded_font_programs(FIXTURE)
            program = next(iter(programs.values()), None)
        requests = [
            FontMetricRequest(word_name="SimSun Probe", family="SimSun", program=program),
            FontMetricRequest(word_name="Latin Probe", family="Times New Roman"),
        ]
        document, entries = build_calibration_document(requests)
        self.assertEqual(
            len(entries),
            len(CALIBRATION_SIZES) * len(requests),
        )
        self.assertTrue(all(entry["page"] == 0 for entry in entries))
        y_values = [entry["y"] for entry in entries]
        self.assertEqual(y_values, sorted(y_values))
        cache_keys = {request.cache_key for request in requests}
        self.assertEqual(len(cache_keys), 2)

    def test_measure_offsets_on_synthetic_pdf(self) -> None:
        from reportlab.pdfgen.canvas import Canvas

        with TemporaryDirectory() as temporary_directory:
            pdf_path = Path(temporary_directory) / "metrics.pdf"
            canvas = Canvas(str(pdf_path), pagesize=(595, 842))
            entries = []
            y = 80.0
            for size in CALIBRATION_SIZES:
                canvas.setFont("Helvetica", size)
                # 让字形框顶部落在 y 附近（drawString 使用基线坐标）
                canvas.drawString(20, 842 - y - 0.72 * size, "probe")
                entries.append(
                    {
                        "word_name": "Probe",
                        "family": "Helvetica",
                        "size": size,
                        "y": y,
                        "page": 0,
                    }
                )
                y += 46.0
            canvas.save()

            samples = measure_offsets(pdf_path, entries)
            self.assertIn("Probe", samples)
            self.assertEqual(len(samples["Probe"]), len(CALIBRATION_SIZES))
            # 同一字体不同字号测得的偏移系数应当接近（用 Helvetica 基线渲染）
            spread = max(samples["Probe"]) - min(samples["Probe"])
            self.assertLess(spread, 0.6)


if __name__ == "__main__":
    unittest.main()
