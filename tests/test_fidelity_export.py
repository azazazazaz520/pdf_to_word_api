from __future__ import annotations

import os
import unittest
import zipfile
from itertools import count
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

from src.validate.report import validate_docx_rendering
from src.ir.builder import build_document_ir
from src.ir.model import IRBlock, IRPage
from src.layout.layout import extract_pdf_layout
from src.export.fidelity import (
    _CanvasContext,
    _FidelityContext,
    _is_irregular_form_table,
    _place_editable_form_table,
    export_fidelity_docx,
)
from src.export.table_render import _set_pdf_cell_content, _table_rows_needing_reflow
from src.layout.models import PdfTable, PdfTableCell


_WORD_AVAILABLE: bool | None = None


def _word_available() -> bool:
    global _WORD_AVAILABLE
    if _WORD_AVAILABLE is not None:
        return _WORD_AVAILABLE
    if os.name != "nt":
        _WORD_AVAILABLE = False
        return _WORD_AVAILABLE
    if os.getenv("PDF_VALIDATION_WORD_RENDER", "1").strip() == "0":
        _WORD_AVAILABLE = False
        return _WORD_AVAILABLE
    import subprocess

    try:
        completed = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-NonInteractive",
                "-Command",
                "$word = New-Object -ComObject Word.Application; "
                "$word.Quit(); Write-Output 'WORD_OK'",
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
            check=False,
        )
        _WORD_AVAILABLE = "WORD_OK" in (completed.stdout or "")
    except Exception:
        _WORD_AVAILABLE = False
    return _WORD_AVAILABLE


def _write_sample_pdf(path: Path, page_count: int = 1, blank_last: bool = False) -> None:
    canvas = Canvas(str(path), pagesize=(595.0, 842.0))
    for page_number in range(1, page_count + 1):
        if blank_last and page_number == page_count:
            canvas.showPage()
            continue
        canvas.setFont("Helvetica", 9)
        canvas.setFillColorRGB(0.2, 0.2, 0.2)
        canvas.drawString(72, 800, "Fidelity Sample Header")
        canvas.drawString(72, 40, f"Page {page_number}")
        canvas.setFillColorRGB(0, 0, 0)
        canvas.setFont("Helvetica-Bold", 16)
        canvas.drawString(72, 760, "1 Introduction")
        canvas.setFont("Helvetica", 11)
        y = 730
        for index in range(5):
            canvas.drawString(
                72,
                y,
                f"Body line {index} used for bbox regression checking.",
            )
            y -= 16
        canvas.setFillColorRGB(0.8, 0.1, 0.1)
        canvas.setFont("Helvetica-Oblique", 12)
        canvas.drawString(200, 620, "Centered red italic caption")
        canvas.setFillColorRGB(0, 0, 0)
        canvas.setStrokeColorRGB(0.1, 0.1, 0.6)
        canvas.setLineWidth(1.2)
        canvas.line(72, 580, 480, 580)
        canvas.setFillColorRGB(0.9, 0.9, 0.95)
        canvas.rect(72, 520, 220, 40, stroke=1, fill=1)
        canvas.showPage()
    canvas.save()


def _build_ir(source_pdf: Path, *, keep_header_footer: bool = True):
    layout = extract_pdf_layout(source_pdf, include_fidelity=True)
    page_sizes = [(page.width, page.height) for page in layout.pages]
    return build_document_ir(
        source_pdf=source_pdf,
        page_routes=["text"] * len(layout.pages),
        page_sizes=page_sizes,
        layout=layout,
        fidelity=True,
        keep_header_footer=keep_header_footer,
        export_mode="fidelity",
    )


def _document_xml(docx_path: Path) -> str:
    with zipfile.ZipFile(docx_path) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def _sections(docx_path: Path) -> list:
    with zipfile.ZipFile(docx_path) as archive:
        root = parse_xml(archive.read("word/document.xml"))
    return list(root.iter(qn("w:sectPr")))


def _frames(docx_path: Path) -> list[dict]:
    with zipfile.ZipFile(docx_path) as archive:
        root = parse_xml(archive.read("word/document.xml"))
    frames: list[dict] = []
    for paragraph in root.iter(qn("w:p")):
        properties = paragraph.find(qn("w:pPr"))
        if properties is None:
            continue
        frame = properties.find(qn("w:framePr"))
        if frame is None:
            continue
        text = "".join(
            node.text or "" for node in paragraph.iter(qn("w:t"))
        )
        sizes = [
            node.get(qn("w:val"))
            for node in paragraph.iter(qn("w:sz"))
            if node.get(qn("w:val"))
        ]
        fonts: list[str] = []
        for node in paragraph.iter(qn("w:rFonts")):
            value = (
                node.get(qn("w:eastAsia"))
                or node.get(qn("w:ascii"))
                or node.get(qn("w:hAnsi"))
            )
            if value:
                fonts.append(value)
        bold_nodes = list(paragraph.iter(qn("w:b")))
        bold = any(
            (node.get(qn("w:val")) or "1") not in {"0", "false"}
            for node in bold_nodes
        )
        frames.append(
            {
                "text": text,
                "x": int(frame.get(qn("w:x")) or 0) / 20.0,
                "y": int(frame.get(qn("w:y")) or 0) / 20.0,
                "width": int(frame.get(qn("w:w")) or 0) / 20.0,
                "font_size": float(sizes[0]) / 2.0 if sizes else 0.0,
                "font": fonts[0] if fonts else "",
                "bold": bold,
            }
        )
    return frames


class FidelityExportTest(unittest.TestCase):
    def test_export_preserves_page_count_and_page_geometry(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            pdf_path = root / "sample.pdf"
            _write_sample_pdf(pdf_path, page_count=2)
            ir = _build_ir(pdf_path)
            output = root / "fidelity.docx"
            export_fidelity_docx(ir, output, source_pdf=pdf_path)

            sections = _sections(output)
            self.assertEqual(len(sections), 2)
            for section, page in zip(sections, ir.pages):
                size = section.find(qn("w:pgSz"))
                self.assertIsNotNone(size)
                self.assertAlmostEqual(
                    int(size.get(qn("w:w"))) / 20.0, page.width, places=1
                )
                self.assertAlmostEqual(
                    int(size.get(qn("w:h"))) / 20.0, page.height, places=1
                )
                margins = section.find(qn("w:pgMar"))
                self.assertEqual(margins.get(qn("w:top")), "0")
                self.assertEqual(margins.get(qn("w:left")), "0")

    def test_text_frames_match_source_bbox_and_font_size(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            pdf_path = root / "sample.pdf"
            _write_sample_pdf(pdf_path, page_count=1)
            layout = extract_pdf_layout(pdf_path, include_fidelity=True)
            ir = _build_ir(pdf_path)
            output = root / "fidelity.docx"
            export_fidelity_docx(ir, output, source_pdf=pdf_path)

            frames = {frame["text"].strip(): frame for frame in _frames(output)}
            body_lines = [
                line
                for line in layout.pages[0].lines
                if not line.is_header_footer
            ]
            self.assertGreaterEqual(len(body_lines), 6)
            checked = 0
            for line in body_lines:
                frame = frames.get(line.text.strip())
                if frame is None:
                    continue
                checked += 1
                self.assertLess(abs(frame["x"] - line.x0), 3.0)
                self.assertLess(abs(frame["y"] - line.top), 3.0)
                self.assertLess(abs(frame["font_size"] - line.font_size), 0.5)
            self.assertGreaterEqual(checked, 6)

    def test_low_confidence_text_remains_editable_without_image_fallback(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            pdf_path = root / "sample.pdf"
            _write_sample_pdf(pdf_path, page_count=1)
            ir = _build_ir(pdf_path)
            target = None
            for block in ir.pages[0].blocks:
                if block.kind in {"paragraph", "heading1"} and block.bbox:
                    target = block
                    break
            self.assertIsNotNone(target)
            target.confidence = 0.2

            hybrid_output = root / "hybrid.docx"
            report = export_fidelity_docx(
                ir,
                hybrid_output,
                source_pdf=pdf_path,
                mode="fidelity_hybrid",
            )
            self.assertEqual(report["fallback_region_count"], 0)
            self.assertEqual(ir.pages[0].fidelity["fallback_regions"], [])
            xml = _document_xml(hybrid_output)
            self.assertIn("Body line 0", xml)
            self.assertNotIn("fallback_1_", xml)
            self.assertFalse(
                any(
                    warning.code == "fidelity_region_fallback"
                    for warning in ir.pages[0].warnings
                )
            )

            strict_ir = _build_ir(pdf_path)
            strict_output = root / "strict.docx"
            strict_report = export_fidelity_docx(
                strict_ir,
                strict_output,
                source_pdf=pdf_path,
                mode="fidelity",
            )
            self.assertEqual(strict_report["fallback_region_count"], 0)

    def test_header_footer_are_written_to_native_parts(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            pdf_path = root / "sample.pdf"
            # 单页文档不会把顶部文本判定为页眉；跨页重复内容才会。
            _write_sample_pdf(pdf_path, page_count=2)
            ir = _build_ir(pdf_path)
            output = root / "fidelity.docx"
            export_fidelity_docx(ir, output, source_pdf=pdf_path)

            with zipfile.ZipFile(output) as archive:
                names = archive.namelist()
                header_names = [name for name in names if "word/header" in name]
                footer_names = [name for name in names if "word/footer" in name]
                self.assertTrue(header_names)
                self.assertTrue(footer_names)
                header_xml = archive.read(header_names[0]).decode("utf-8")
                footer_xml = archive.read(footer_names[0]).decode("utf-8")
            self.assertIn("Fidelity Sample Header", header_xml)
            self.assertIn("Page 1", footer_xml)
            self.assertTrue(ir.pages[0].header_footer_native)

    def test_source_blank_page_is_preserved(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            pdf_path = root / "blank-last.pdf"
            _write_sample_pdf(pdf_path, page_count=2, blank_last=True)
            ir = _build_ir(pdf_path)
            self.assertEqual(ir.pages[1].route, "blank")
            output = root / "fidelity.docx"
            export_fidelity_docx(ir, output, source_pdf=pdf_path)
            self.assertEqual(len(_sections(output)), 2)

    def test_vectors_are_exported_as_absolute_shapes(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            pdf_path = root / "sample.pdf"
            _write_sample_pdf(pdf_path, page_count=1)
            ir = _build_ir(pdf_path)
            self.assertGreaterEqual(ir.vector_count, 2)
            output = root / "fidelity.docx"
            export_fidelity_docx(ir, output, source_pdf=pdf_path)
            xml = _document_xml(output)
            self.assertIn("<wps:wsp>", xml)
            self.assertIn("1A1A99", xml.upper())
            self.assertEqual(ir.pages[0].fidelity["native_block_count"] > 0, True)

    def test_embedded_image_keeps_source_bbox(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            image_path = root / "logo.png"
            pdf_path = root / "image.pdf"
            Image.new("RGB", (120, 60), (30, 120, 200)).save(image_path)
            canvas = Canvas(str(pdf_path), pagesize=(595.0, 842.0))
            canvas.drawImage(
                ImageReader(str(image_path)), 400, 700, width=120, height=60
            )
            canvas.showPage()
            canvas.save()

            ir = _build_ir(pdf_path)
            output = root / "fidelity.docx"
            export_fidelity_docx(ir, output, source_pdf=pdf_path)
            xml = _document_xml(output)
            self.assertIn("<wp:anchor", xml)
            image_block = next(
                block
                for block in ir.pages[0].blocks
                if block.kind == "image"
            )
            self.assertAlmostEqual(image_block.bbox[0], 400.0, places=1)
            self.assertAlmostEqual(image_block.bbox[1], 82.0, places=1)

    def test_bordered_table_uses_floating_fixed_layout(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            pdf_path = root / "table.pdf"
            canvas = Canvas(str(pdf_path), pagesize=(595.0, 842.0))
            rows = (
                ("ID", "Product", "Qty", "Amount"),
                ("T-001", "OCR text", "12", "1,234.56"),
                ("T-002", "Table structure", "8", "800.00"),
            )
            column_x = (72.0, 160.0, 340.0, 410.0)
            for row_index, values in enumerate(rows):
                y = 700 - row_index * 30
                for x, value in zip(column_x, values):
                    canvas.drawString(x + 6, y, value)
            canvas.rect(72, 640, 460, 90, stroke=1, fill=0)
            for x in column_x:
                canvas.line(x, 640, x, 730)
            for row_index in range(4):
                canvas.line(72, 640 + row_index * 30, 532, 640 + row_index * 30)
            canvas.save()

            layout = extract_pdf_layout(pdf_path, include_fidelity=True)
            self.assertEqual(len(layout.pages[0].tables), 1)
            ir = _build_ir(pdf_path)
            self.assertEqual(ir.table_count, 1)
            output = root / "fidelity.docx"
            export_fidelity_docx(ir, output, source_pdf=pdf_path)
            xml = _document_xml(output)
            self.assertIn("<w:tblpPr", xml)
            self.assertIn('w:tblLayout w:type="fixed"', xml)
            self.assertIn('w:hRule="exact"', xml)
            self.assertEqual(len(_sections(output)), 1)

    def test_table_font_shrink_compares_pdf_points_with_cell_points(self) -> None:
        document = Document()
        table = document.add_table(rows=1, cols=2)

        def write_cell(cell, right_edge: float) -> float:
            glyphs = tuple(
                SimpleNamespace(
                    text=character,
                    bbox=(index * right_edge / 6.0, 0.0, (index + 1) * right_edge / 6.0, 10.0),
                )
                for index, character in enumerate("Longer")
            )
            span = SimpleNamespace(
                glyphs=glyphs,
                font_name="Arial",
                font_size=10.0,
                bold=False,
                italic=False,
                color=None,
            )
            _set_pdf_cell_content(
                cell,
                "Longer",
                width=2.0,
                is_header=False,
                spans=(span,),
                shrink_cell_font=True,
            )
            run = next(run for run in cell.paragraphs[0].runs if run.text == "Longer")
            return run.font.size.pt

        self.assertAlmostEqual(write_cell(table.cell(0, 0), 90.0), 10.0)
        self.assertAlmostEqual(write_cell(table.cell(0, 1), 130.0), 9.0)

    def test_table_reflow_keeps_narrow_replacement_font_text_visible(self) -> None:
        table = PdfTable(
            bbox=(0.0, 0.0, 252.0, 21.0),
            column_boundaries=(0.0, 58.0, 252.0),
            row_boundaries=(0.0, 21.0),
            rows=(("Permanently and totally disabled", "Table structure"),),
            cells=(
                PdfTableCell(
                    row_index=0,
                    column_index=0,
                    row_span=1,
                    column_span=1,
                    bbox=(0.0, 0.0, 58.0, 21.0),
                    text="Permanently and totally disabled",
                    font_size=7.0,
                ),
                PdfTableCell(
                    row_index=0,
                    column_index=1,
                    row_span=1,
                    column_span=1,
                    bbox=(58.0, 0.0, 252.0, 21.0),
                    text="Table structure",
                    font_size=7.0,
                ),
            ),
        )

        self.assertEqual(
            _table_rows_needing_reflow(
                table,
                row_offset=0,
                cell_padding_points=(2.0, 0.0),
            ),
            {0},
        )

    def test_only_irregular_merged_form_tables_are_selected_for_local_fallback(self) -> None:
        regular = PdfTable(
            bbox=(0.0, 0.0, 180.0, 24.0),
            column_boundaries=(0.0, 60.0, 120.0, 180.0),
            row_boundaries=(0.0, 12.0, 24.0),
            rows=(("A", "B", "C"), ("1", "2", "3")),
            cells=tuple(
                PdfTableCell(
                    row_index=row,
                    column_index=column,
                    row_span=1,
                    column_span=1,
                    bbox=(column * 60.0, row * 12.0, (column + 1) * 60.0, (row + 1) * 12.0),
                    text=value,
                    font_size=8.0,
                )
                for row, values in enumerate((("A", "B", "C"), ("1", "2", "3")))
                for column, value in enumerate(values)
            ),
        )
        irregular = PdfTable(
            bbox=(0.0, 0.0, 180.0, 60.0),
            column_boundaries=(0.0, 60.0, 120.0, 180.0),
            row_boundaries=(0.0, 30.0, 60.0),
            rows=(("Long form field text", "", ""), ("", "", "")),
            cells=(
                PdfTableCell(
                    row_index=0,
                    column_index=0,
                    row_span=2,
                    column_span=2,
                    bbox=(0.0, 0.0, 120.0, 60.0),
                    text="Long form field text that needs a local form treatment and an additional label",
                    font_size=8.0,
                ),
            ),
        )
        dotted_form = PdfTable(
            bbox=(0.0, 0.0, 180.0, 120.0),
            column_boundaries=(0.0, 60.0, 120.0, 180.0),
            row_boundaries=tuple(float(index * 10) for index in range(13)),
            rows=tuple(
                (". . .", "", str(index))
                for index in range(12)
            ),
            cells=tuple(
                PdfTableCell(
                    row_index=index,
                    column_index=0,
                    row_span=1,
                    column_span=1,
                    bbox=(0.0, index * 10.0, 60.0, (index + 1) * 10.0),
                    text=". . .",
                    font_size=8.0,
                )
                for index in range(6)
            ),
        )

        self.assertFalse(_is_irregular_form_table(regular))
        self.assertTrue(_is_irregular_form_table(irregular))
        self.assertTrue(_is_irregular_form_table(dotted_form))

        short_form = PdfTable(
            bbox=(0.0, 0.0, 240.0, 24.0),
            column_boundaries=tuple(float(index * 30) for index in range(9)),
            row_boundaries=(0.0, 12.0, 24.0),
            rows=(("Long form label with enough descriptive field text", "", "", "", "", "", ""), ("", "", "", "", "", "", "", "")),
            cells=tuple(
                PdfTableCell(
                    row_index=0,
                    column_index=index,
                    row_span=1,
                    column_span=1,
                    bbox=(index * 30.0, 0.0, (index + 1) * 30.0, 12.0),
                    text="Long form label with enough descriptive field text" if index == 0 else "",
                    font_size=8.0,
                )
                for index in range(10)
            ),
        )
        self.assertTrue(_is_irregular_form_table(short_form))

    def test_irregular_form_table_rebuild_keeps_text_editable(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        context = _FidelityContext(
            document=document,
            source_pdf=None,
            mode="fidelity",
            min_confidence=0.0,
            fallback_dpi=200.0,
            fallback_max_pixels=100000,
            page_width_points=200.0,
            object_ids=count(1),
        )
        span = SimpleNamespace(
            text="Editable field",
            bbox=(12.0, 14.0, 72.0, 22.0),
            glyphs=(),
            font_name="Arial",
            pdf_font_name="Helvetica",
            font_size=8.0,
            color=(0, 0, 0),
            bold=False,
            italic=False,
        )
        cell = SimpleNamespace(bbox=(10.0, 10.0, 90.0, 30.0), spans=(span,))
        table = SimpleNamespace(
            cells=(cell,),
            has_borders=True,
            border_color=(0, 0, 0),
            border_width=0.0,
        )
        page = IRPage(page_number=1, width=200.0, height=300.0, route="text")
        block = IRBlock(
            kind="table",
            page=1,
            bbox=(10.0, 10.0, 90.0, 30.0),
            table=table,
            z_order=2,
        )
        report = {"placements": []}

        rebuilt = _place_editable_form_table(
            context,
            _CanvasContext(document=document, paragraph=paragraph),
            page,
            block,
            report=report,
        )

        self.assertTrue(rebuilt)
        self.assertEqual(report["editable_form_table_count"], 1)
        self.assertEqual(report["editable_form_text_count"], 1)
        self.assertEqual(report["placements"][0]["reason"], "positioned_form_table")
        xml = document.element.xml
        self.assertIn("Editable field", xml)
        self.assertIn("form_table_edge_1_", xml)

    def test_mixed_font_line_exports_separate_font_frames(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            pdf_path = root / "mixed-font.pdf"
            canvas = Canvas(str(pdf_path), pagesize=(400, 200))
            canvas.setFont("Helvetica", 12)
            canvas.drawString(40, 150, "Plain")
            canvas.setFont("Helvetica-Bold", 12)
            canvas.drawString(120, 150, "BOLD123")
            canvas.save()

            layout = extract_pdf_layout(pdf_path, include_fidelity=True)
            ir = build_document_ir(
                source_pdf=pdf_path,
                page_routes=["text"],
                page_sizes=[(layout.pages[0].width, layout.pages[0].height)],
                layout=layout,
                fidelity=True,
                keep_header_footer=True,
                export_mode="fidelity",
            )
            output = root / "fidelity.docx"
            export_fidelity_docx(ir, output, source_pdf=pdf_path)

            frames = _frames(output)
            texts = {frame["text"]: frame for frame in frames}
            self.assertIn("Plain", texts)
            self.assertIn("BOLD123", texts)
            self.assertFalse(texts["Plain"]["bold"])
            self.assertTrue(texts["BOLD123"]["bold"])
            self.assertAlmostEqual(
                texts["BOLD123"]["x"],
                texts["Plain"]["x"] + 80.0,
                delta=3.0,
            )

    @unittest.skipUnless(_word_available(), "缺少 Microsoft Word，跳过渲染回读验收")
    def test_word_render_meets_acceptance_criteria(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            pdf_path = root / "sample.pdf"
            _write_sample_pdf(pdf_path, page_count=1)
            ir = _build_ir(pdf_path)
            output = root / "fidelity.docx"
            export_fidelity_docx(ir, output, source_pdf=pdf_path)

            result = validate_docx_rendering(
                output,
                source_page_count=len(ir.pages),
                work_dir=root,
                timeout_seconds=240.0,
                source_pdf=pdf_path,
                compare_ssim=True,
                compare_text=True,
                ssim_threshold=0.98,
            )
            self.assertEqual(result["page_delta"], 0)
            self.assertEqual(result["unexpected_blank_pages"], [])
            ssim = result["ssim"]
            self.assertEqual(ssim["status"], "succeeded")
            self.assertGreaterEqual(ssim["min_ssim"], 0.98)
            layout = result["text_layout"]
            self.assertEqual(layout["status"], "succeeded")
            self.assertLess(layout["max_bbox_error"], 3.0)
            self.assertLess(layout["max_font_size_error"], 0.5)


if __name__ == "__main__":
    unittest.main()
