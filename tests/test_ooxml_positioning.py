from __future__ import annotations

import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

from src.ooxml_positioning import (
    add_absolute_picture,
    add_absolute_shape,
    add_absolute_text_box,
    add_absolute_text_paragraph,
    header_footer_paragraph,
    points_to_emu,
    points_to_twips,
    position_table,
    prepare_header,
    rgb_to_hex,
    set_exact_page,
)


class OoxmlPositioningTest(unittest.TestCase):
    def test_unit_conversions(self) -> None:
        self.assertEqual(points_to_emu(1.0), 12700)
        self.assertEqual(points_to_emu(72.0), 914400)
        self.assertEqual(points_to_twips(1.0), 20)
        self.assertEqual(points_to_twips(72.0), 1440)
        self.assertEqual(rgb_to_hex((255, 128, 0)), "FF8000")
        self.assertEqual(rgb_to_hex(None, "112233"), "112233")

    def test_exact_page_size_uses_source_dimensions(self) -> None:
        document = Document()
        section = document.sections[0]
        info = set_exact_page(section, width_points=595.0, height_points=842.0)
        self.assertFalse(info["scaled"])
        self.assertAlmostEqual(section.page_width.pt, 595.0, places=2)
        self.assertAlmostEqual(section.page_height.pt, 842.0, places=2)
        self.assertEqual(section.left_margin.pt, 0.0)

    def test_oversized_page_is_scaled_with_flag(self) -> None:
        document = Document()
        info = set_exact_page(
            document.sections[0], width_points=2000.0, height_points=3000.0
        )
        self.assertTrue(info["scaled"])
        self.assertAlmostEqual(info["height_points"], 22.0 * 72.0, places=2)

    def test_frame_paragraph_is_absolutely_positioned(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "frame.docx"
            document = Document()
            set_exact_page(
                document.sections[0], width_points=595.0, height_points=842.0
            )
            add_absolute_text_paragraph(
                document,
                "Absolutely positioned text",
                x_points=100.0,
                y_points=200.0,
                width_points=200.0,
                font_name="Arial",
                font_size=12.0,
                bold=True,
                color=(200, 30, 30),
            )
            document.save(output)
            with __import__("zipfile").ZipFile(output) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("<w:framePr", xml)
            self.assertIn('<w:sz w:val="24"/>', xml)  # 12pt -> 24 half points
            self.assertIn('w:x="2000"', xml)  # 100pt -> 2000 twips
            self.assertIn('w:y="4000"', xml)
            self.assertIn('w:hAnchor="page"', xml)
            self.assertIn('w:vAnchor="page"', xml)
            self.assertIn("C81E1E", xml)

    def test_anchored_picture_uses_page_offsets(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "picture.docx"
            document = Document()
            set_exact_page(
                document.sections[0], width_points=595.0, height_points=842.0
            )
            buffer = BytesIO()
            Image.new("RGB", (40, 20), (10, 120, 200)).save(buffer, format="PNG")
            add_absolute_picture(
                document,
                buffer.getvalue(),
                x_points=120.0,
                y_points=340.0,
                width_points=40.0,
                height_points=20.0,
            )
            document.save(output)
            with __import__("zipfile").ZipFile(output) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("<wp:anchor", xml)
            self.assertIn("<wp:posOffset>1524000</wp:posOffset>", xml)  # 120pt
            self.assertIn("<wp:posOffset>4318000</wp:posOffset>", xml)  # 340pt

    def test_absolute_shape_and_table_positioning(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "shapes.docx"
            document = Document()
            set_exact_page(
                document.sections[0], width_points=595.0, height_points=842.0
            )
            add_absolute_shape(
                document,
                geometry="line",
                x_points=72.0,
                y_points=500.0,
                width_points=300.0,
                height_points=1.5,
                line_color=(20, 20, 200),
                line_width_points=1.5,
            )
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "A"
            table.cell(0, 1).text = "B"
            position_table(table, x_points=72.0, y_points=600.0)
            document.save(output)
            with __import__("zipfile").ZipFile(output) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("<wps:wsp>", xml)
            self.assertIn("<w:tblpPr", xml)
            self.assertIn('w:tblpX="1440"', xml)
            self.assertIn('w:tblpY="12000"', xml)

    def test_text_box_contains_txbx_content(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "textbox.docx"
            document = Document()
            set_exact_page(
                document.sections[0], width_points=595.0, height_points=842.0
            )
            add_absolute_text_box(
                document,
                x_points=100.0,
                y_points=120.0,
                width_points=200.0,
                height_points=30.0,
                text="Rotated text",
                rotation=15.0,
            )
            document.save(output)
            with __import__("zipfile").ZipFile(output) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            self.assertIn("<wps:txbx>", xml)
            self.assertIn("Rotated text", xml)
            self.assertIn('rot="900000"', xml)

    def test_header_paragraph_can_be_prepared(self) -> None:
        document = Document()
        header = prepare_header(document.sections[0])
        paragraph = header_footer_paragraph(header)
        add_absolute_text_paragraph(
            header,
            "Header text",
            x_points=72.0,
            y_points=30.0,
            width_points=200.0,
            font_size=9.0,
            paragraph=paragraph,
        )
        self.assertEqual(paragraph.text, "Header text")
        self.assertFalse(header.is_linked_to_previous)


if __name__ == "__main__":
    unittest.main()
